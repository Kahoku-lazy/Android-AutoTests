"""File upload endpoints — avatars and chat attachments."""

import base64
import json
import logging
import uuid

from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt

from ..decorators import require_auth
from ..upload_cleanup import ensure_upload_dir, remove_upload_file

logger = logging.getLogger("ai_assistant")

UPLOAD_DIR = ensure_upload_dir()

DOCUMENT_EXTENSIONS = {
    "txt",
    "log",
    "md",
    "markdown",
    "json",
    "xml",
    "csv",
    "py",
    "js",
    "html",
    "css",
    "yaml",
    "yml",
    "docx",
    "xlsx",
    "pdf",
}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif"}
ALLOWED_EXTENSIONS = DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
MAX_UPLOAD_SIZE = 20 * 1024 * 1024  # 20MB documents
MAX_IMAGE_UPLOAD_SIZE = 5 * 1024 * 1024  # 5MB images


def _image_media_type(ext: str) -> str:
    return "image/jpeg" if ext == "jpg" else f"image/{ext}"


def _parse_docx(filepath: str) -> str:
    try:
        from docx import Document

        doc = Document(filepath)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        raise ValueError(f"DOCX 解析失败: {e}") from e


def _parse_xlsx(filepath: str) -> str:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(filepath, data_only=True)
        parts = []
        for name in wb.sheetnames:
            ws = wb[name]
            parts.append(f"--- Sheet: {name} ---")
            for row in ws.iter_rows(values_only=True):
                parts.append("\t".join(str(c) if c is not None else "" for c in row))
        return "\n".join(parts)
    except Exception as e:
        raise ValueError(f"XLSX 解析失败: {e}") from e


def _parse_pdf(filepath: str) -> str:
    try:
        import fitz

        doc = fitz.open(filepath)
        parts = []
        for page in doc:
            text = page.get_text()
            if text.strip():
                parts.append(f"--- Page {page.number + 1} ---\n{text}")
        return "\n".join(parts) if parts else "[PDF has no extractable text]"
    except Exception as e:
        raise ValueError(f"PDF 解析失败: {e}") from e


def _parse_markdown(filepath: str) -> str:
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            return f.read()
    except Exception as e:
        raise ValueError(f"MD 解析失败: {e}") from e


@csrf_exempt
@require_auth
def upload_avatar(request):
    """Upload an avatar image and return a data URI stored directly in the DB.

    Avatars are stored as data URIs (data:image/png;base64,...) in the
    AIAgent.avatar field so they travel with the database — no disk files
    that can get lost between environments.
    """
    try:
        data = json.loads(request.body)
        img_b64 = data.get("image", "")
        if not img_b64:
            return JsonResponse({"status": False, "message": "no image data"})
        # Normalise: strip the data: prefix if present, so we can prepend our own
        if "," in img_b64:
            img_b64 = img_b64.split(",", 1)[1]
        # Validate that it's actually decodable base64
        base64.b64decode(img_b64)
        data_uri = f"data:image/png;base64,{img_b64}"
        return JsonResponse({"status": True, "url": data_uri})
    except Exception:
        logger.exception("Avatar upload failed")
        return JsonResponse({"status": False, "message": "头像上传失败，请重试"}, status=500)


@csrf_exempt
@require_auth
def upload_and_parse_file(request):
    if request.method != "POST":
        return JsonResponse({"status": False, "message": "POST required"}, status=405)

    uploaded = request.FILES.get("file")
    if not uploaded:
        return JsonResponse({"status": False, "message": "No file uploaded"}, status=400)

    ext = uploaded.name.rsplit(".", 1)[-1].lower() if "." in uploaded.name else ""
    if ext not in ALLOWED_EXTENSIONS:
        return JsonResponse(
            {
                "status": False,
                "message": f"Unsupported file type: .{ext}. Supported: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
            },
            status=400,
        )

    max_size = MAX_IMAGE_UPLOAD_SIZE if ext in IMAGE_EXTENSIONS else MAX_UPLOAD_SIZE
    if uploaded.size > max_size:
        return JsonResponse(
            {
                "status": False,
                "message": f"File too large ({uploaded.size} bytes). Max: {max_size} bytes",
            },
            status=400,
        )

    saved_name = f"{uuid.uuid4().hex}_{uploaded.name}"
    filepath = UPLOAD_DIR / saved_name
    try:
        with open(filepath, "wb") as f:
            for chunk in uploaded.chunks():
                f.write(chunk)

        if ext in IMAGE_EXTENSIONS:
            raw = filepath.read_bytes()
            media_type = _image_media_type(ext)
            data_b64 = base64.b64encode(raw).decode("ascii")
            return JsonResponse(
                {
                    "status": True,
                    "data": {
                        "filename": uploaded.name,
                        "size": uploaded.size,
                        "type": ext,
                        "media_type": media_type,
                        "data_uri": f"data:{media_type};base64,{data_b64}",
                    },
                }
            )

        text_extensions = {
            "txt",
            "log",
            "json",
            "xml",
            "csv",
            "py",
            "js",
            "html",
            "css",
            "yaml",
            "yml",
        }
        content = ""
        parse_error = None

        try:
            if ext in text_extensions:
                content = open(filepath, "r", encoding="utf-8", errors="replace").read()
            elif ext in ("md", "markdown"):
                content = _parse_markdown(str(filepath))
            elif ext == "docx":
                content = _parse_docx(str(filepath))
            elif ext == "xlsx":
                content = _parse_xlsx(str(filepath))
            elif ext == "pdf":
                content = _parse_pdf(str(filepath))
        except Exception as e:
            parse_error = str(e)
            content = f"[Parse error: {e}]"

        if len(content) > 50000:
            content = content[:50000] + f"\n\n[... truncated {len(content) - 50000} chars]"

        return JsonResponse(
            {
                "status": True,
                "data": {
                    "filename": uploaded.name,
                    "size": uploaded.size,
                    "type": ext,
                    "content": content,
                    "preview": content[:300] + ("..." if len(content) > 300 else ""),
                    "parse_error": parse_error,
                },
            }
        )
    finally:
        remove_upload_file(filepath)
