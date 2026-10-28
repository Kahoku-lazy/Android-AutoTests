"""任务附件解析 — 配置 YAML 原样保留，文档转为 Markdown。"""

from __future__ import annotations

import uuid

from dataclasses import dataclass
from pathlib import Path

from utils.logs import backend_log

DOCUMENT_EXTENSIONS = {".md", ".docx", ".pdf"}
CONFIG_EXTENSIONS = {".yaml", ".yml"}
TASK_ATTACHMENT_MAX_SIZE = 20 * 1024 * 1024
MAX_CASE_DOCUMENTS = 5


class AttachmentError(ValueError):
    """用户可见的任务附件错误。"""


def _suffix(filename: str) -> str:
    return Path(str(filename or "")).suffix.lower()


def _safe_name(filename: str) -> str:
    raw = str(filename or "").replace("\\", "/").strip()
    base = Path(raw).name.strip()
    if not base or base in {".", ".."}:
        raise AttachmentError("非法文件名")
    return base


def _check_size(content: bytes) -> None:
    if len(content) > TASK_ATTACHMENT_MAX_SIZE:
        raise AttachmentError("文件过大")


def _fallback_docx(path: str) -> str:
    """python-docx 兜底解析：保留段落、标题级别与表格。"""
    from docx import Document

    document = Document(path)
    parts: list[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = str(getattr(getattr(para, "style", None), "name", "") or "").lower()
        if style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(int(digits or 1), 6)
            parts.append(f"{'#' * level} {text}")
        else:
            parts.append(text)
    for table in document.tables:
        rows = [
            [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            for row in table.rows
        ]
        rows = [row for row in rows if any(row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        for row in rows:
            row.extend([""] * (width - len(row)))
        parts.append("| " + " | ".join(rows[0]) + " |")
        parts.append("|" + "---|" * width)
        for row in rows[1:]:
            parts.append("| " + " | ".join(row) + " |")
    return "\n\n".join(parts).strip()


def _fallback_pdf(path: str) -> str:
    """PyMuPDF 兜底解析：按页提取可复制文本。"""
    import fitz

    document = fitz.open(path)
    parts: list[str] = []
    for page in document:
        text = (page.get_text() or "").strip()
        if text:
            parts.append(f"--- Page {page.number + 1} ---\n{text}")
    return "\n\n".join(parts).strip()


def _fallback_convert(path: str) -> str:
    """按扩展名走本地解析；没有可用兜底时返回空串。"""
    ext = _suffix(Path(path).name)
    if ext == ".docx":
        return _fallback_docx(path)
    if ext == ".pdf":
        return _fallback_pdf(path)
    return ""


def convert_with_markitdown(path: str) -> str:
    """把 Word / PDF 转成 Markdown；MarkItDown 缺可选依赖时自动走本地兜底。

    Args:
        path: 临时文件路径。

    Returns:
        解析出的 Markdown 文本。

    Raises:
        AttachmentError: MarkItDown 与本地兜底都无法解析，或没有可用文本。
    """
    from markitdown import MarkItDown

    markitdown_error: Exception | None = None
    try:
        result = MarkItDown().convert(path)
        text = str(getattr(result, "text_content", "") or "").strip()
        if text:
            return text
    except Exception as exc:
        markitdown_error = exc
        backend_log.warning("MarkItDown 解析失败，改用本地兜底: %s: %s", type(exc).__name__, exc)

    try:
        fallback = _fallback_convert(path)
    except Exception as exc:
        backend_log.exception("本地文档解析失败: %s", path)
        raise AttachmentError(f"文档解析失败: {exc}") from exc
    if fallback:
        return fallback
    if markitdown_error is not None:
        raise AttachmentError(f"文档解析失败: {markitdown_error}") from markitdown_error
    raise AttachmentError("文档解析失败: 没有可用文本")


def parse_document_bytes(filename: str, content: bytes) -> tuple[str, str]:
    """文档附件 → (markdown, 安全文件名)。失败抛 AttachmentError。"""
    _check_size(content)
    base = _safe_name(filename)
    ext = _suffix(base)
    if ext not in DOCUMENT_EXTENSIONS:
        raise AttachmentError("文档仅支持 Markdown、Word（.docx）与 PDF")
    if ext == ".md":
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise AttachmentError("Markdown 必须是 UTF-8") from exc
        return text, base
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name
    try:
        try:
            text = convert_with_markitdown(tmp_path)
        except AttachmentError:
            raise
        except Exception as exc:
            backend_log.exception("任务附件解析失败: %s", base)
            raise AttachmentError(f"文档解析失败: {exc}") from None
    finally:
        Path(tmp_path).unlink(missing_ok=True)
    return text, base


def parse_task_attachment_bytes(filename: str, content: bytes) -> tuple[str, str]:
    """兼容旧调用名，行为与 parse_document_bytes 相同。"""
    return parse_document_bytes(filename, content)


@dataclass(frozen=True)
class PreparedUpload:
    kind: str
    filename: str
    body: str
    content: bytes


def _upload_pair(item) -> tuple[str, bytes]:
    if hasattr(item, "read"):
        return _safe_name(getattr(item, "name", "") or "file"), item.read()
    filename, content = item
    return _safe_name(filename), content


def collect_task_files(configs: list, documents: list) -> list[PreparedUpload]:
    """检查一次提交的配置与文档。失败抛 AttachmentError。不写库、不落盘。"""
    from engines.agents.workflow.case_file import validate_platform_yaml

    if not configs:
        raise AttachmentError("需要一份 YAML 配置")
    if len(configs) != 1:
        raise AttachmentError("只能有一份配置")
    filename, content = _upload_pair(configs[0])
    _check_size(content)
    ext = _suffix(filename)
    if ext not in CONFIG_EXTENSIONS:
        raise AttachmentError("配置只支持 YAML（.yaml / .yml）")
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise AttachmentError("YAML 必须是 UTF-8") from exc
    _cases, reason = validate_platform_yaml(text)
    if reason:
        raise AttachmentError(reason)
    prepared = [PreparedUpload(kind="config", filename=filename, body=text, content=content)]
    for item in documents:
        doc_name, doc_bytes = _upload_pair(item)
        markdown, safe = parse_document_bytes(doc_name, doc_bytes)
        prepared.append(
            PreparedUpload(kind="document", filename=safe, body=markdown, content=doc_bytes)
        )
    return prepared


def store_flat_upload(filename: str, content: bytes) -> str:
    """把原件写成上传目录第一层的文件，返回相对文件名。"""
    from backend.apps.task_list.upload_cleanup import ensure_upload_dir

    root = ensure_upload_dir()
    target = root / f"{uuid.uuid4().hex}_{_safe_name(filename)}"
    target.write_bytes(content)
    return target.name


def collect_document_files(documents: list) -> list[PreparedUpload]:
    """检查用例编写任务上传的 PRD/Word 文档。失败抛 AttachmentError，不写库、不落盘。

    Args:
        documents: 上传文件列表（Django UploadedFile 或 (文件名, 字节) 二元组）。

    Returns:
        解析后的文档列表，kind 固定为 document。

    Raises:
        AttachmentError: 没有文档、文档数量超限，或某份文档解析失败。
    """
    if not documents:
        raise AttachmentError("请上传 PRD 或 Word 文档")
    if len(documents) > MAX_CASE_DOCUMENTS:
        raise AttachmentError(f"一次最多上传 {MAX_CASE_DOCUMENTS} 份文档")
    prepared: list[PreparedUpload] = []
    for item in documents:
        doc_name, doc_bytes = _upload_pair(item)
        markdown, safe = parse_document_bytes(doc_name, doc_bytes)
        prepared.append(
            PreparedUpload(kind="document", filename=safe, body=markdown, content=doc_bytes)
        )
    return prepared
