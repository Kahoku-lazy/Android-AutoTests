#!/usr/bin/env python3
"""生成 PPT 1：项目开发工作流全景 (from agent-workflow-report.html)"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── 颜色常量 ──
PRIMARY = RGBColor(0x19, 0xC8, 0xB9)
TEXT = RGBColor(0x79, 0x4F, 0x27)
TEXT_BODY = RGBColor(0x72, 0x5D, 0x42)
TEXT_SEC = RGBColor(0x9F, 0x92, 0x7D)
TEXT_MUTED = RGBColor(0x8A, 0x7B, 0x66)
BG = RGBColor(0xF8, 0xF8, 0xF0)
BG_CONTENT = RGBColor(0xF7, 0xF3, 0xDF)
BORDER = RGBColor(0xC4, 0xB8, 0x9E)
SUCCESS = RGBColor(0x6F, 0xBA, 0x2C)
WARNING = RGBColor(0xF5, 0xC3, 0x1C)
ERROR = RGBColor(0xE0, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PHASE_PLAN = RGBColor(0xF7, 0xCD, 0x67)
PHASE_DESIGN = RGBColor(0x88, 0x9D, 0xF0)
PHASE_DEV = RGBColor(0x6F, 0xBA, 0x2C)
PHASE_TEST = RGBColor(0xB3, 0x9E, 0xF3)
PHASE_RELEASE = RGBColor(0xF8, 0xA6, 0xB2)
PHASE_DELIVER = RGBColor(0x19, 0xC8, 0xB9)

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)

SAVE_DIR = os.path.dirname(os.path.abspath(__file__))


def add_bg(slide, color=BG):
    """设置幻灯片背景色"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=14, color=TEXT_BODY,
                 bold=False, alignment=PP_ALIGN.LEFT, font_name='Noto Sans SC'):
    """添加文本框"""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_rounded_rect(slide, left, top, width, height, fill_color=BG_CONTENT, border_color=BORDER):
    """添加圆角矩形"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    return shape


def add_rect(slide, left, top, width, height, fill_color=BG_CONTENT):
    """添加矩形"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_slide_number(slide, num):
    """添加页码"""
    add_text_box(slide, 15, 8.4, 0.8, 0.4, str(num), font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.RIGHT)


# ═══════════════════════ Slide 1: 标题页 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide, BG)
add_rounded_rect(slide, 1.5, 1.5, 13, 5.5, BG_CONTENT, BORDER)

add_text_box(slide, 2, 2.2, 12, 1.5, '项目开发工作流全景', font_size=44, color=TEXT, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 2, 3.8, 12, 0.8, '从立项到交付的完整生命周期 · PRD 驱动 · ARCH 约束 · CHECKLIST 验收',
             font_size=18, color=TEXT_SEC, alignment=PP_ALIGN.CENTER)

# KPI 小标签
tags = [('立项', PHASE_PLAN), ('需求与设计', PHASE_DESIGN), ('开发', PHASE_DEV),
        ('发布与评估', PHASE_RELEASE), ('交付', PHASE_DELIVER)]
for i, (label, color) in enumerate(tags):
    x = 3.5 + i * 2
    shape = add_rounded_rect(slide, x, 5.2, 1.7, 0.5, WHITE, color)
    shape.text_frame.paragraphs[0].text = f'● {label}'
    shape.text_frame.paragraphs[0].font.size = Pt(11)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT_BODY
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 2, 6.5, 12, 0.5, 'Android-AutoTests · 2026-07-16 · v2.0', font_size=14, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ═══════════════════════ Slide 2: 项目全景 KPI ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '📊 项目全景 — 关键数据', font_size=28, color=TEXT, bold=True)

kpis = [('5', '项目阶段', '立项→设计→开发→发布→交付'),
        ('3', '文档体系', 'PRD · ARCH · CHECKLIST'),
        ('267', '验收项 (PRD)', '覆盖 8 个模块的完整验收条件'),
        ('28', '检查项 (CHECKLIST)', '6 组自检守门员'),
        ('4', '变更级别', '🔴大改 🟡增量 🟢修补 ⚪重构')]

for i, (num, label, desc) in enumerate(kpis):
    x = 0.5 + i * 3.1
    add_rounded_rect(slide, x, 1.5, 2.8, 2.5, BG_CONTENT, BORDER)
    add_text_box(slide, x + 0.2, 1.7, 2.4, 1.0, num, font_size=48, color=TEXT, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.2, 2.7, 2.4, 0.5, label, font_size=16, color=TEXT_SEC, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.2, 3.3, 2.4, 0.6, desc, font_size=11, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# 五大阶段
add_text_box(slide, 0.5, 4.5, 15, 0.6, '📋 五大阶段流程', font_size=22, color=TEXT, bold=True)
phases = [('第一阶段\n立项', PHASE_PLAN, '产品定位·竞品分析\n灵感分析·价值评估'),
          ('第二阶段\n需求与设计', PHASE_DESIGN, 'PRD输出·ARCH设计\n需求评审·架构评审'),
          ('第三阶段\n项目开发', PHASE_DEV, '任务拆分·需求迭代\n功能开发·测试验证'),
          ('第四阶段\n发布与评估', PHASE_RELEASE, '全功能验收·风险评估\n安全·性能·易用性'),
          ('第五阶段\n产品交付', PHASE_DELIVER, '部署文档·用户文档\n回滚预案·监控告警')]
for i, (name, color, desc) in enumerate(phases):
    x = 0.5 + i * 3.1
    shape = add_rounded_rect(slide, x, 5.3, 2.8, 3.0, BG_CONTENT, color)
    add_text_box(slide, x + 0.1, 5.4, 2.6, 1.0, name, font_size=15, color=TEXT, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 6.6, 2.6, 1.5, desc, font_size=11, color=TEXT_BODY, alignment=PP_ALIGN.CENTER)

add_slide_number(slide, 2)

# ═══════════════════════ Slide 3: 第一阶段 - 立项 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '一、项目初识阶段 — 第一阶段：立项', font_size=28, color=TEXT, bold=True)

items = [('1.1 产品定位', '明确产品定位与用户', '为什么要做这个产品？面向哪类用户？哪些客户有需求？输出一句话定位 + 目标用户画像。'),
         ('1.2 竞品分析', '研究竞品技术、设计与思路', '研究同类项目的技术栈、架构设计、做得好的与不好的地方。'),
         ('1.3 灵感分析', '从多项目寻找设计灵感', '分析多种同类型项目，找设计灵感。思考如何提升用户体验，改善产品交互与功能性。'),
         ('1.4 价值评估', '通过竞品市场与收益评估价值', '竞品有多少用户？盈利模式是什么？这个赛道是否值得投入？输出立项验证报告。')]

for i, (num, name, desc) in enumerate(items):
    y = 1.5 + i * 1.8
    add_rounded_rect(slide, 0.5, y, 15, 1.5, BG_CONTENT, PHASE_PLAN)
    add_text_box(slide, 0.8, y + 0.15, 2, 0.4, num, font_size=13, color=RGBColor(0xD4, 0xA0, 0x30), bold=True)
    add_text_box(slide, 0.8, y + 0.5, 14, 0.4, name, font_size=17, color=TEXT, bold=True)
    add_text_box(slide, 0.8, y + 0.9, 14, 0.5, desc, font_size=12, color=TEXT_BODY)

add_text_box(slide, 0.5, 8.5, 15, 0.4, '📁 产出物：dev_docs/01-立项与设计/立项验证/', font_size=11, color=TEXT_MUTED)
add_slide_number(slide, 3)

# ═══════════════════════ Slide 4: 第二阶段 - 需求与设计 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '一、项目初识阶段 — 第二阶段：需求与设计', font_size=28, color=TEXT, bold=True)

items2 = [('2.1 产品需求', '输出完整的产品需求文档',
           '需求大纲：项目定位·用户画像·业务逻辑架构·模块交互场景\n子模块 PRD：功能定位·页面交互·数据表单·异常场景·功能边界\n每个功能点附带验收条件 + 测试方案'),
          ('2.2 架构设计', '输出项目架构设计',
           '架构大纲：三层分离·模块依赖·交互边界·部署·设计决策\n子模块 ARCH：组件树·API设计·数据模型·交互时序·跨模块通信\n技术选型、命名标准、UI 设计素材'),
          ('2.3 需求评审', '评审需求文档',
           '需求是否覆盖所有用户场景？验收条件是否可测试？功能边界是否明确？异常场景是否完整？'),
          ('2.4 架构评审', '评审架构设计',
           '架构三层边界是否清晰？模块间交互是否合理？技术选型是否合适？数据模型是否完整？')]

for i, (num, name, desc) in enumerate(items2):
    y = 1.5 + i * 1.8
    add_rounded_rect(slide, 0.5, y, 15, 1.5, BG_CONTENT, PHASE_DESIGN)
    add_text_box(slide, 0.8, y + 0.15, 2, 0.4, num, font_size=13, color=RGBColor(0x40, 0x60, 0xC0), bold=True)
    add_text_box(slide, 0.8, y + 0.5, 14, 0.4, name, font_size=17, color=TEXT, bold=True)
    add_text_box(slide, 0.8, y + 0.9, 14, 0.5, desc, font_size=11, color=TEXT_BODY)

add_text_box(slide, 0.5, 8.5, 15, 0.4, '📁 产出物：dev_docs/02-PRD需求/ + dev_docs/03-设计与架构/', font_size=11, color=TEXT_MUTED)
add_slide_number(slide, 4)

# ═══════════════════════ Slide 5: 第三阶段 - 开发 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '一、项目初识阶段 — 第三阶段：项目开发', font_size=28, color=TEXT, bold=True)

# 开发流程小标签
dev_flow = [('📋 任务拆分', PHASE_PLAN), ('📝 需求迭代', PHASE_DESIGN), ('💻 功能开发', PHASE_DEV),
            ('🧪 测试验证', PHASE_TEST), ('✅ 评审', PHASE_RELEASE), ('🔄 复盘', PHASE_DELIVER)]
for i, (label, color) in enumerate(dev_flow):
    x = 0.5 + i * 2.55
    shape = add_rounded_rect(slide, x, 1.3, 2.3, 0.6, RGBColor(0xFF, 0xFF, 0xFF), color)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

dev_items = [
    ('3.1 任务拆分', '根据需求与架构拆分任务', PHASE_PLAN,
     '按模块拆分开发任务，明确每个任务的输入（PRD功能块）、输出（代码+自测结果）、依赖关系。'),
    ('3.2 需求迭代', '需求端负责需求与功能设计迭代', PHASE_DESIGN,
     '开发过程中发现需求不明确→回PRD补充验收条件。发现功能边界冲突→回需求大纲调整边界定义。'),
    ('3.3 功能开发', '开发端按需求实现功能', PHASE_DEV,
     '代码不偏离架构设计→对照ARCH文档的文件结构。功能正确实现→对照PRD验收条件。写完自测→对照CHECKLIST。'),
    ('3.4 测试验证', '测试端根据需求制定方案并验证', PHASE_TEST,
     '测试方案来源=PRD中的测试场景（前置→操作→期望）。输出：测试结果+checklist表单。'),
    ('3.5 三方评审', '需求·开发·测试三方信息同步', PHASE_RELEASE,
     '对照进度→确保文档与代码同步→测试结果与开发进度同步。三方中任何一方发现不一致→标记为阻塞项。'),
    ('3.6 问题复盘', '高频问题、验证问题复盘', PHASE_DELIVER,
     '记录：现象→根因→解决方案→测试方案。同类问题≥3次→新增规则到CHECKLIST或.claude/rules/。'),
]
for i, (num, name, color, desc) in enumerate(dev_items):
    row = i // 2
    col = i % 2
    x = 0.5 + col * 7.7
    y = 2.2 + row * 2.2
    add_rounded_rect(slide, x, y, 7.4, 1.9, BG_CONTENT, color)
    add_text_box(slide, x + 0.3, y + 0.1, 7, 0.3, num, font_size=12, color=TEXT_MUTED, bold=True)
    add_text_box(slide, x + 0.3, y + 0.45, 7, 0.35, name, font_size=15, color=TEXT, bold=True)
    add_text_box(slide, x + 0.3, y + 0.85, 7, 0.9, desc, font_size=11, color=TEXT_BODY)

add_slide_number(slide, 5)

# ═══════════════════════ Slide 6: 第四阶段 - 发布与风险评估 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '一、项目初识阶段 — 第四阶段：发布与风险评估', font_size=28, color=TEXT, bold=True)

rel_items = [('4.1 全功能验收', '交付前对照需求验收一遍',
              '以PRD验收条件为基准，逐功能走一遍完整的用户操作流程。核心链路（设备→元素→用例→执行→报告）必须全部通过。'),
             ('4.2 风险评估', '评估项目发布风险',
              '哪些功能是新增的（高风险）？哪些是修改的（中风险）？哪些未改动（低风险）？有没有跨模块改动（高风险）？'),
             ('4.3 安全·性能·稳定性', '非功能需求检查',
              'JWT鉴权是否正常？API Key是否正确加密？有没有SQL注入风险？大页面加载是否超过3秒？并发锁定是否安全？'),
             ('4.4 易用性检查', '考虑生产环境的易用性问题',
              '首次使用是否有引导？错误提示是否用户能看懂（无技术术语）？没有设备时是否有清晰的操作指引？')]

for i, (num, name, desc) in enumerate(rel_items):
    y = 1.5 + i * 1.8
    add_rounded_rect(slide, 0.5, y, 15, 1.5, BG_CONTENT, PHASE_RELEASE)
    add_text_box(slide, 0.8, y + 0.15, 2, 0.4, num, font_size=13, color=RGBColor(0xE0, 0x68, 0x80), bold=True)
    add_text_box(slide, 0.8, y + 0.5, 14, 0.4, name, font_size=17, color=TEXT, bold=True)
    add_text_box(slide, 0.8, y + 0.9, 14, 0.5, desc, font_size=12, color=TEXT_BODY)

add_slide_number(slide, 6)

# ═══════════════════════ Slide 7: 第五阶段 - 交付 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '一、项目初识阶段 — 第五阶段：产品交付', font_size=28, color=TEXT, bold=True)

del_items = [('5.1 部署文档', '部署步骤与依赖清单',
              '环境要求（OS/Python/Node/ADB）、依赖安装（pip/npm）、数据库初始化、Redis配置、一键启动脚本验证。'),
             ('5.2 用户文档', '快速上手指南',
              '从登录到完成第一次测试的完整操作指南。含截图和常见问题解答。'),
             ('5.3 回滚预案', '发布失败的回滚方案',
              '数据库migration回滚步骤、上一版本备份位置、回滚后验证清单。'),
             ('5.4 监控与告警', '上线后的监控指标',
              '关键API错误率、设备连接成功率、执行引擎队列积压、前端页面加载时间。异常时通知谁？')]

for i, (num, name, desc) in enumerate(del_items):
    y = 1.5 + i * 1.8
    add_rounded_rect(slide, 0.5, y, 15, 1.5, BG_CONTENT, PHASE_DELIVER)
    add_text_box(slide, 0.8, y + 0.15, 2, 0.4, num, font_size=13, color=PRIMARY, bold=True)
    add_text_box(slide, 0.8, y + 0.5, 14, 0.4, name, font_size=17, color=TEXT, bold=True)
    add_text_box(slide, 0.8, y + 0.9, 14, 0.5, desc, font_size=12, color=TEXT_BODY)

add_slide_number(slide, 7)

# ═══════════════════════ Slide 8: 项目迭代 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '二、项目迭代 — 已有需求与架构后的执行', font_size=28, color=TEXT, bold=True)

# 变更需求流程
add_text_box(slide, 0.5, 1.3, 15, 0.5, '变更需求流程', font_size=20, color=TEXT, bold=True)
change_flow = [('🔴/🟡/🟢\n判定级别', PHASE_PLAN), ('先改 PRD\n写验收条件', PHASE_DESIGN),
               ('改 ARCH\n(如需要)', PHASE_TEST), ('写代码\n对照验收', PHASE_DEV), ('跑 CHECKLIST\n通过→合并', PHASE_DELIVER)]
for i, (label, color) in enumerate(change_flow):
    x = 0.5 + i * 3.1
    shape = add_rounded_rect(slide, x, 2.0, 2.8, 1.2, WHITE, color)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(13)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 0.5, 3.5, 15, 0.4, '核心原则：先改文档，后写代码。PRD是开发的输入来源，不是开发完的回忆录。',
             font_size=13, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# 代码实现 - 变更级别
add_text_box(slide, 0.5, 4.2, 15, 0.5, '代码实现 — 按变更级别执行', font_size=20, color=TEXT, bold=True)
impl_items = [('变更级别', '🔴 大改：PRD→ARCH→代码全走一遍\n🟡 增量：PRD加验收项→代码→跑CHECKLIST\n🟢 修补：改代码→检查PRD受影响项\n⚪ 重构：改代码→更新ARCH文件路径'),
              ('Bug修复', '四层下钻定位：L1浏览器→L2前端编译→L3 API→L4数据库\n对照PRD确认"预期行为"vs"实际行为"差异\n对照ARCH检查架构边界是否被破坏')]
for i, (title, desc) in enumerate(impl_items):
    x = 0.5 + i * 7.7
    add_rounded_rect(slide, x, 4.9, 7.4, 1.8, BG_CONTENT, PHASE_DEV)
    add_text_box(slide, x + 0.3, 5.0, 7, 0.35, title, font_size=15, color=TEXT, bold=True)
    add_text_box(slide, x + 0.3, 5.4, 7, 1.2, desc, font_size=11, color=TEXT_BODY)

# 测试保障
add_text_box(slide, 0.5, 7.0, 15, 0.4, '测试保障 — 六大检查项', font_size=20, color=TEXT, bold=True)
checks = ['① 功能有没有问题？', '② 是不是按需求做的？', '③ 架构有没有偏离？',
          '④ 改了A，B坏了没？', '⑤ 核心链路通不通？', '⑥ 改了文档没有？']
for i, check in enumerate(checks):
    x = 0.5 + i * 2.55
    shape = add_rounded_rect(slide, x, 7.6, 2.3, 0.6, BG_CONTENT, BORDER)
    shape.text_frame.paragraphs[0].text = check
    shape.text_frame.paragraphs[0].font.size = Pt(10)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT_BODY
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_slide_number(slide, 8)

# ═══════════════════════ Slide 9: AI 协作 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '三、AI 协作 — 我提需求，你执行，我验收', font_size=28, color=TEXT, bold=True)

# AI 工作流
add_text_box(slide, 0.5, 1.3, 15, 0.5, '你提需求 → AI 的工作流', font_size=20, color=TEXT, bold=True)
ai_flow = [('👤 你\n提出需求', PHASE_PLAN), ('📖 阅读 PRD\n理解上下文', PHASE_DESIGN),
           ('📐 对照 ARCH\n设计方案', PHASE_TEST), ('💻 执行\n写代码', PHASE_DEV),
           ('🔍 自检\n跑 CHECKLIST', PHASE_RELEASE), ('📋 输出结果\n等你验收', PHASE_DELIVER)]
for i, (label, color) in enumerate(ai_flow):
    x = 0.3 + i * 2.6
    shape = add_rounded_rect(slide, x, 2.0, 2.3, 1.2, WHITE, color)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# AI 使用的文档
add_text_box(slide, 0.5, 3.5, 15, 0.5, 'AI 使用的项目文档与工具体系', font_size=20, color=TEXT, bold=True)
docs = [('📋 PRD 需求文档', 'AI理解"做什么"—267条验收条件+101个测试场景'),
        ('🏗️ ARCH 架构文档', 'AI理解"怎么做"—组件树、API、数据模型、交互时序'),
        ('✅ CHECKLIST', 'AI自检"做对没有"—6组28条检查项'),
        ('🔧 gen_arch_stats.py', 'AI自动检测漂移—表数、端点数、Tool数、步骤类型'),
        ('📖 技术栈参考', '命名统一标准、UI设计素材、API契约'),
        ('🤖 .claude/rules/', '行为准则、历史高频问题速查、模块边界防火墙')]
for i, (name, desc) in enumerate(docs):
    row = i // 3
    col = i % 3
    x = 0.5 + col * 5.1
    y = 4.2 + row * 1.2
    add_rounded_rect(slide, x, y, 4.8, 1.0, BG_CONTENT, BORDER)
    add_text_box(slide, x + 0.2, y + 0.1, 4.4, 0.35, name, font_size=13, color=TEXT, bold=True)
    add_text_box(slide, x + 0.2, y + 0.5, 4.4, 0.45, desc, font_size=10, color=TEXT_MUTED)

# 验收措施
add_text_box(slide, 0.5, 6.8, 15, 0.4, '验收措施 — 确保AI按你的想法执行', font_size=20, color=TEXT, bold=True)
measures = ['① 文档先行', '② CHECKLIST守门', '③ 自动检测不漂移',
            '④ 变更级别判定', '⑤ 先方案后执行', '⑥ 结果对照PRD']
for i, m in enumerate(measures):
    x = 0.5 + i * 2.55
    shape = add_rounded_rect(slide, x, 7.4, 2.3, 0.6, BG_CONTENT, PHASE_RELEASE)
    shape.text_frame.paragraphs[0].text = m
    shape.text_frame.paragraphs[0].font.size = Pt(11)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT_BODY
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_slide_number(slide, 9)

# ═══════════════════════ Slide 10: 文档-代码-验收 铁三角 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '四、文档 · 代码 · 验收 铁三角', font_size=28, color=TEXT, bold=True)

# 铁三角循环
tri_items = [('📋 PRD 需求', PHASE_PLAN, '267 验收项\n功能定位与边界\n异常场景与容错'),
             ('🏗️ ARCH 架构', PHASE_DESIGN, '组件树 + API\n数据模型 + 交互时序\n自动统计附录'),
             ('✅ CHECKLIST', PHASE_DEV, '6 组 28 条\n4 级变更判定\n快速自检 + 完整检查'),
             ('💻 代码', PHASE_TEST, 'apps/ 8 模块\nfrontend/ 8 模块\nagentscope_service/'),
             ('🚀 交付', PHASE_DELIVER, '验收通过\n文档同步\n可追溯')]
for i, (name, color, desc) in enumerate(tri_items):
    x = 0.5 + i * 3.1
    shape = add_rounded_rect(slide, x, 1.5, 2.8, 2.8, color, BORDER)
    add_text_box(slide, x + 0.1, 1.7, 2.6, 0.7, name, font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 2.7, 2.6, 1.3, desc, font_size=11, color=WHITE, alignment=PP_ALIGN.CENTER)
    if i < 4:
        add_text_box(slide, x + 2.85, 2.5, 0.4, 0.5, '⇄', font_size=24, color=TEXT_SEC, bold=True)

# 关键文件索引
add_text_box(slide, 0.5, 4.8, 15, 0.5, '关键文件索引', font_size=20, color=TEXT, bold=True)
files = [
    ('dev_docs/02-PRD需求/需求大纲.md', '项目定位·用户画像·模块职责边界', '需求变更时更新'),
    ('dev_docs/02-PRD需求/PRD-0X-*.md', '8模块功能需求+验收条件+测试方案', '功能变更时更新'),
    ('dev_docs/03-设计与架构/架构大纲.md', '三层分离·模块依赖图·设计决策', '架构变更时更新'),
    ('dev_docs/03-设计与架构/ARCH-0X-*.md', '8模块技术架构（组件树/API/数据模型）', '代码结构变更时更新'),
    ('dev_docs/DEVELOPMENT_CHECKLIST.md', '验收守门员—6组28条+4级变更判定', '发现问题模式时更新'),
    ('tools/gen_arch_stats.py', '自动统计（表/端点/Tool/步骤类型）', '新增统计维度时更新'),
]
for i, (path, role, maintainer) in enumerate(files):
    y = 5.5 + i * 0.55
    add_text_box(slide, 0.8, y, 5.5, 0.45, path, font_size=10, color=PRIMARY, bold=True)
    add_text_box(slide, 6.5, y, 5.5, 0.45, role, font_size=10, color=TEXT_BODY)
    add_text_box(slide, 12.2, y, 3.5, 0.45, maintainer, font_size=10, color=TEXT_MUTED)

add_slide_number(slide, 10)

# ── 保存 ──
output_path = os.path.join(SAVE_DIR, '01-项目开发工作流全景.pptx')
prs.save(output_path)
print(f'✅ PPT 1 已保存: {output_path}')
