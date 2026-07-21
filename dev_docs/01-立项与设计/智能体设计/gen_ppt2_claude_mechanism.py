#!/usr/bin/env python3
"""生成 PPT 2：Claude Code 运行机制 (from 智能体体系设计-迭代版本-20260716.html)"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# ── 颜色常量 ──
PRIMARY = RGBColor(0x19, 0xC8, 0xB9)
TEXT = RGBColor(0x79, 0x4F, 0x27)
TEXT_BODY = RGBColor(0x72, 0x5D, 0x42)
TEXT_SEC = RGBColor(0x9F, 0x92, 0x7D)
TEXT_MUTED = RGBColor(0x8A, 0x7B, 0x66)
BG = RGBColor(0xF8, 0xF8, 0xF0)
BG_CONTENT = RGBColor(0xF7, 0xF3, 0xDF)
BORDER = RGBColor(0xC4, 0xB8, 0x9E)
SUCCESS = RGBColor(0x6F, 0xBA, 0x2C)
WARNING = RGBColor(0xF5, 0xC3, 0x1C)
ERROR = RGBColor(0xE0, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLAUDE = RGBColor(0xF7, 0xCD, 0x67)

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)

SAVE_DIR = os.path.dirname(os.path.abspath(__file__))


def add_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=14, color=TEXT_BODY,
                 bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = 'Noto Sans SC'
    p.alignment = alignment
    return tf


def add_rounded_rect(slide, left, top, width, height, fill_color=BG_CONTENT, border_color=BORDER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    return shape


def add_slide_number(slide, num):
    add_text_box(slide, 15, 8.4, 0.8, 0.4, str(num), font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.RIGHT)


# ═══════════════════════ Slide 1: 标题页 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_rounded_rect(slide, 1.5, 1.5, 13, 5.5, BG_CONTENT, BORDER)

add_text_box(slide, 2, 2.2, 12, 1.5, 'Claude Code 运行机制', font_size=44, color=TEXT, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 2, 3.6, 12, 0.8, 'v2 · 已补全 CLAUDE.md + 7 Agent + 14 Rules + 12 Skills · 全部 git 追踪',
             font_size=18, color=TEXT_SEC, alignment=PP_ALIGN.CENTER)

# KPI tags
tags = [('CLAUDE.md', CLAUDE), ('7 Agents', RGBColor(0x88, 0x9D, 0xF0)),
        ('14 Rules', SUCCESS), ('12 Skills', RGBColor(0xB3, 0x9E, 0xF3)), ('55 Memories', PRIMARY)]
for i, (label, color) in enumerate(tags):
    x = 3 + i * 2.1
    shape = add_rounded_rect(slide, x, 5.2, 1.8, 0.5, WHITE, color)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT_BODY
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 2, 6.5, 12, 0.5, 'Android-AutoTests · 2026-07-16', font_size=14, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ═══════════════════════ Slide 2: 启动时得到什么 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '一、启动时我得到了什么', font_size=28, color=TEXT, bold=True)

# 系统注入上下文
add_text_box(slide, 0.5, 1.3, 7.5, 0.5, '系统注入的上下文（每次对话自动获得）', font_size=18, color=TEXT, bold=True)
sys_items = [('Git 状态 + 最近 commits', '知道当前分支、未提交改动、最近变更方向'),
             ('当前工作目录', '所有相对路径基于此解析'),
             ('平台信息 (macOS + zsh)', '决定shell语法、路径分隔符'),
             ('Memory 文件索引', 'MEMORY.md加载→知道沉淀的教训、偏好、约定'),
             ('可用工具列表', 'Bash/Read/Write/Edit/Agent/Skill/MCP等'),
             ('可用 Agent/Skill 列表', '.claude/agents/*.md + .claude/skills/')]
for i, (src, usage) in enumerate(sys_items):
    y = 2.0 + i * 0.65
    add_rounded_rect(slide, 0.5, y, 7.5, 0.55, BG_CONTENT, BORDER)
    add_text_box(slide, 0.8, y + 0.08, 3, 0.4, src, font_size=12, color=TEXT, bold=True)
    add_text_box(slide, 4.0, y + 0.08, 3.8, 0.4, usage, font_size=11, color=TEXT_MUTED)

# 项目实际情况
add_text_box(slide, 8.5, 1.3, 7, 0.5, '项目实际情况（v2 · 2026-07-16）', font_size=18, color=TEXT, bold=True)
configs = [('CLAUDE.md', '已创建', '路由表·身份·铁律'),
           ('.claude/agents/', '7 个', 'architect/developer/prd-writer等'),
           ('.claude/rules/', '14 个', 'frontend/backend/architecture等'),
           ('.claude/hooks/', '6 个', 'check-boundary/doc-drift等'),
           ('.claude/skills/', '12 个', 'auto-dev/prd-writer等'),
           ('.claude/memory/', '~55 个', '教训沉淀+MEMORY.md索引'),
           ('dev_docs/', '已追踪', 'PRD+ARCH+CHECKLIST')]
for i, (item, status, desc) in enumerate(configs):
    y = 2.0 + i * 0.65
    add_rounded_rect(slide, 8.5, y, 7, 0.55, BG_CONTENT, BORDER)
    tag_color = SUCCESS
    add_text_box(slide, 8.8, y + 0.08, 2.5, 0.4, item, font_size=11, color=TEXT, bold=True)
    add_text_box(slide, 11.5, y + 0.08, 1.2, 0.4, f'✅ {status}', font_size=10, color=SUCCESS, bold=True)
    add_text_box(slide, 13.0, y + 0.08, 2.3, 0.4, desc, font_size=9, color=TEXT_MUTED)

add_text_box(slide, 8.5, 6.8, 7, 0.6, 'v1（旧版）时 CLAUDE.md、AGENTS.md、.claude/rules/ 均不存在，dev_docs/ 在 gitignore 中。\nv2 全部补齐并纳入 git 追踪。',
             font_size=10, color=TEXT_MUTED)
add_slide_number(slide, 2)

# ═══════════════════════ Slide 3: 信息获取优先级 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '二、我如何从项目中获取信息', font_size=28, color=TEXT, bold=True)

# 三层优先级
layers = [
    ('第 1 层：上下文已有的', CLAUDE, [
        'Memory 文件索引（MEMORY.md 加载时带入）— 包含相关教训/偏好',
        'System prompt 中的项目信息（git status、最近 commits、可用工具）'
    ]),
    ('第 2 层：按需读取的', PRIMARY, [
        'dev_docs/ — 涉及功能设计→Read PRD；涉及架构→Read ARCH',
        '源代码 — Glob/Find 定位→Read 打开→LSP 跳转定义/引用',
        'gen_arch_stats.py — 需要数字时运行（表数/端点数/Tool数/步骤类型）'
    ]),
    ('第 3 层：需要对比时才读的', RGBColor(0x88, 0x9D, 0xF0), [
        'PRD vs 代码 — 验证功能是否按需求实现',
        'ARCH vs 代码 — gen_arch_stats.py --check-md 自动对比',
        'DEVELOPMENT_CHECKLIST — 开发完成后逐项自检（6组28条）'
    ]),
]
for i, (title, color, items) in enumerate(layers):
    y = 1.5 + i * 2.4
    add_rounded_rect(slide, 0.5, y, 15, 2.1, BG_CONTENT, color)
    add_text_box(slide, 0.8, y + 0.1, 14, 0.4, title, font_size=16, color=TEXT, bold=True)
    for j, item in enumerate(items):
        add_text_box(slide, 1.2, y + 0.6 + j * 0.45, 13.5, 0.4, f'• {item}', font_size=12, color=TEXT_BODY)

# 文档→代码映射
add_text_box(slide, 0.5, 7.8, 15, 0.3, '文档 → 代码映射：PRD验收条件→前端index.vue+后端views.py | ARCH API表→urls.py | ARCH数据模型→models.py',
             font_size=11, color=TEXT_MUTED)
add_slide_number(slide, 3)

# ═══════════════════════ Slide 4: 处理流程 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '三、你提问后我的处理流程', font_size=28, color=TEXT, bold=True)

# 完整链路
flow_steps = [('收到消息', CLAUDE), ('判断类型', RGBColor(0x88, 0x9D, 0xF0)), ('确定范围', SUCCESS),
              ('获取信息', PRIMARY), ('推理/执行', RGBColor(0xB3, 0x9E, 0xF3)), ('输出结果', WARNING)]
for i, (label, color) in enumerate(flow_steps):
    x = 0.5 + i * 2.6
    shape = add_rounded_rect(slide, x, 1.5, 2.3, 0.8, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(14)
    shape.text_frame.paragraphs[0].font.color.rgb = WHITE if color not in [CLAUDE, WARNING] else TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# Step 1 & 2
add_text_box(slide, 0.5, 2.7, 7.5, 0.4, 'Step 1: 判断消息类型', font_size=16, color=TEXT, bold=True)
types = ['A. 纯问答：「这个模块做什么的？」→ 只读，直接读文档回答',
         'B. 代码修改：「改一下这个功能」→ 评估影响面、改代码、验证',
         'C. 文档维护：「更新一下 PRD」→ 只改文档，不改代码',
         'D. 运行验证：「跑一下看看」→ 启动应用，操作验证，输出结果']
for i, t in enumerate(types):
    add_text_box(slide, 0.8, 3.2 + i * 0.42, 7, 0.38, t, font_size=11, color=TEXT_BODY)

add_text_box(slide, 8.5, 2.7, 7, 0.4, 'Step 2: 确定范围和级别', font_size=16, color=TEXT, bold=True)
levels = ['🔴 大改：新增模块/改枚举/改架构 → 先读PRD→ARCH→出方案→等确认→执行',
          '🟡 增量：加功能/加API → 先读PRD→加验收条件→改代码→跑CHECKLIST',
          '🟢 修补：改文案/修CSS → 读PRD确认受影响项→改代码→检查',
          '⚪ 重构：搬家/改名 → 改代码→更新ARCH路径→跑gen_arch_stats.py']
for i, l in enumerate(levels):
    add_text_box(slide, 8.8, 3.2 + i * 0.42, 6.8, 0.38, l, font_size=11, color=TEXT_BODY)

# Step 3-6
add_text_box(slide, 0.5, 5.2, 15, 0.4, 'Step 3-6: 以代码修改为例的执行细节', font_size=16, color=TEXT, bold=True)
steps_3_6 = [
    ('3. 获取信息', '读PRD→读ARCH→Grep/Glob定位文件→Read确认当前内容'),
    ('4. 推理/执行', '🔴大改→EnterPlanMode出计划等批 | 🟡增量→直接编辑+gen_arch_stats | 🟢修补→一次性改完'),
    ('5. 自检', '对照CHECKLIST逐项确认→PRD验收条件vs改动行为→gen_arch_stats通过？→跨模块越界？'),
    ('6. 输出', '说明改了什么+为什么这样改+对照了哪些PRD验收条件+gen_arch_stats结果'),
]
for i, (title, desc) in enumerate(steps_3_6):
    y = 5.8 + i * 0.7
    add_rounded_rect(slide, 0.5, y, 15, 0.6, BG_CONTENT, BORDER)
    add_text_box(slide, 0.8, y + 0.12, 2.5, 0.35, title, font_size=13, color=TEXT, bold=True)
    add_text_box(slide, 3.5, y + 0.12, 11.5, 0.35, desc, font_size=11, color=TEXT_BODY)

add_slide_number(slide, 4)

# ═══════════════════════ Slide 5: 文档维护逻辑 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '四、我是如何维护项目文档的', font_size=28, color=TEXT, bold=True)

# 决策表
add_text_box(slide, 0.5, 1.3, 15, 0.5, '文档维护决策树', font_size=20, color=TEXT, bold=True)

# 表头
headers = ['代码变更', 'PRD更新？', 'ARCH更新？', 'CHECKLIST更新？']
for i, h in enumerate(headers):
    x = 0.5 + i * 3.8
    add_rounded_rect(slide, x, 2.0, 3.5, 0.5, CLAUDE, BORDER)
    shape = add_rounded_rect(slide, x, 2.0, 3.5, 0.5, CLAUDE, BORDER)
    shape.text_frame.paragraphs[0].text = h
    shape.text_frame.paragraphs[0].font.size = Pt(14)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

rows = [
    ('新增模块', '✅ 新增子PRD+需求大纲索引', '✅ 新增ARCH+架构大纲依赖图', '✅ 新模块跨模块边界规则'),
    ('新增API端点', '✅ 对应功能块验收条件', '✅ §4 API表+1行', '❌ 除非引入新类型问题'),
    ('新增数据表/字段', '✅ §3 数据表单约束', '✅ §5 数据模型ER图', '❌'),
    ('新增前端组件', '✅ §2 交互描述', '✅ §2 组件树+1节点', '❌'),
    ('改StepType枚举', '✅ 所有涉及步骤数的章节', '✅ 架构大纲§九', '✅ gen_arch_stats自动同步'),
    ('改错误提示文案', '✅ 对应异常场景提示词', '❌', '❌'),
    ('文件移动/重命名', '❌', '✅ 所有路径引用', '❌'),
    ('修CSS/样式', '❌', '❌', '❌'),
]
for i, row in enumerate(rows):
    y = 2.6 + i * 0.55
    for j, cell in enumerate(row):
        x = 0.5 + j * 3.8
        add_rounded_rect(slide, x, y, 3.5, 0.5, BG_CONTENT, BORDER)
        add_text_box(slide, x + 0.15, y + 0.08, 3.2, 0.35, cell, font_size=11, color=TEXT_BODY if j > 0 else TEXT, bold=(j == 0))

# 主动读文档规则
add_text_box(slide, 0.5, 7.2, 15, 0.4, '主动读文档规则', font_size=20, color=TEXT, bold=True)
rules = ['🔴 大改动：先读PRD（理解功能意图）→ 再读ARCH（理解技术约束）→ 再读代码（看现有实现）→ 三个都读完才动手',
         '🟡 增量：先读PRD对应功能块（确认验收条件）→ 读ARCH对应模块（确认文件位置）→ 改代码',
         '🟢 修补：直接读代码改 → 改完对比PRD异常场景确认文案一致 → ARCH通常不需要动']
for i, r in enumerate(rules):
    add_text_box(slide, 0.8, 7.7 + i * 0.38, 14.5, 0.35, r, font_size=11, color=TEXT_BODY)

add_slide_number(slide, 5)

# ═══════════════════════ Slide 6: 编码思考逻辑 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '五、编写代码时我如何思考', font_size=28, color=TEXT, bold=True)

# 思考链路
add_text_box(slide, 0.5, 1.3, 15, 0.5, '思考链路（以"给元素表格加一个筛选条件"为例）', font_size=18, color=TEXT, bold=True)
think_steps = [
    ('① 定位', '前端在哪？→ frontend/src/modules/element-locator/ → ElementManager.vue\n后端在哪？→ apps/element_locator/ → views.py'),
    ('② 参照', '已有的筛选Tab怎么写的？→ Grep "filter"或"Tab" → 读已有代码的DOM结构和事件处理 → 用同样的模式'),
    ('③ 模仿', '复制已有筛选Tab的结构 → 改label/筛选逻辑 → 编码风格完全对齐 → 不引入新的写法'),
    ('④ 验证', '改完后：→ PRD验收条件打勾 → gen_arch_stats.py通过 → CHECKLIST组1/2/4通过 → 浏览器操作确认'),
]
for i, (title, desc) in enumerate(think_steps):
    x = 0.5 + i * 3.85
    add_rounded_rect(slide, x, 2.0, 3.55, 2.8, BG_CONTENT, BORDER)
    add_text_box(slide, x + 0.2, 2.2, 3, 0.4, title, font_size=16, color=TEXT, bold=True)
    add_text_box(slide, x + 0.2, 2.8, 3.2, 1.8, desc, font_size=11, color=TEXT_BODY)

# 代码风格约束
add_text_box(slide, 0.5, 5.2, 15, 0.4, '代码风格约束', font_size=18, color=TEXT, bold=True)
add_text_box(slide, 0.5, 5.8, 7.5, 0.35, '✅ 我会做的', font_size=15, color=SUCCESS, bold=True)
do_items = ['模仿已有代码的命名、注释密度、缩进风格', '新增API用 {ok, data|error} 统一格式',
            '前端组件用 PascalCase.vue，JS变量 camelCase', '数据库表用正确前缀（el_/dp_/cm_...）',
            '读PRD异常场景确认错误提示文案', '改完跑 gen_arch_stats.py']
for i, item in enumerate(do_items):
    add_text_box(slide, 0.8, 6.3 + i * 0.38, 7, 0.35, f'→ {item}', font_size=11, color=TEXT_BODY)

add_text_box(slide, 8.5, 5.8, 7, 0.35, '❌ 我不做的', font_size=15, color=ERROR, bold=True)
dont_items = ['前端直连数据库', '跨App import内部实现（service/runner/consumer）',
              '仪表盘中写 INSERT/UPDATE/DELETE', '错误提示暴露技术术语给用户',
              '硬编码 API Key 或密码', '静默吞异常（catch必须log或报错）']
for i, item in enumerate(dont_items):
    add_text_box(slide, 8.8, 6.3 + i * 0.38, 7, 0.35, f'→ {item}', font_size=11, color=TEXT_BODY)

add_slide_number(slide, 6)

# ═══════════════════════ Slide 7: 问题排查 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '六、代码出问题时我如何思考和解决', font_size=28, color=TEXT, bold=True)

# 四层下钻
add_text_box(slide, 0.5, 1.3, 15, 0.5, '四层下钻排查（从用户症状到根因）', font_size=20, color=TEXT, bold=True)
drill = [('L1 浏览器', WARNING, '页面白屏/报错/按钮没反应', 'DevTools Console/Network\nVue组件渲染'),
         ('L2 前端编译', SUCCESS, '编译失败/组件不渲染', 'npx vite build\nVue SFC语法/import路径'),
         ('L3 API', RGBColor(0x88, 0x9D, 0xF0), 'API 500/404/数据不对', 'urls.py路由\nviews.py逻辑+curl验证'),
         ('L4 数据库', PRIMARY, '数据不一致/丢失', 'models.py约束\nmigrations+ORM查询')]
for i, (level, color, symptom, tool) in enumerate(drill):
    x = 0.5 + i * 3.85
    add_rounded_rect(slide, x, 2.0, 3.55, 2.5, BG_CONTENT, color)
    add_text_box(slide, x + 0.2, 2.2, 3, 0.4, level, font_size=16, color=TEXT, bold=True)
    add_text_box(slide, x + 0.2, 2.7, 3, 0.4, f'症状：{symptom}', font_size=11, color=TEXT_BODY)
    add_text_box(slide, x + 0.2, 3.3, 3, 0.6, f'工具：\n{tool}', font_size=11, color=TEXT_MUTED)

# 反思闭环
add_text_box(slide, 0.5, 4.9, 15, 0.4, '修复后的反思闭环', font_size=20, color=TEXT, bold=True)
reflections = [
    ('第 1 次出现', '记录到本次会话的问题清单。', SUCCESS),
    ('第 2 次出现', '⚠ 提醒你"这个问题之前出现过，建议加规则"。', WARNING),
    ('第 3+ 次出现', '🔴 写入 .claude/memory/ + 建议更新 CHECKLIST。同类错误不允许第4次。', ERROR),
]
for i, (title, desc, color) in enumerate(reflections):
    x = 0.5 + i * 5.1
    add_rounded_rect(slide, x, 5.5, 4.8, 1.2, BG_CONTENT, color)
    add_text_box(slide, x + 0.2, 5.7, 4.4, 0.35, title, font_size=15, color=TEXT, bold=True)
    add_text_box(slide, x + 0.2, 6.1, 4.4, 0.5, desc, font_size=11, color=TEXT_BODY)

# 排查流程
add_text_box(slide, 0.5, 7.0, 15, 0.3, '排查流程：用户反馈 → L1浏览器 → L2前端编译 → L3 API → L4数据库 → 修复+规则沉淀',
             font_size=12, color=TEXT_MUTED)
add_slide_number(slide, 7)

# ═══════════════════════ Slide 8: 上下文管理 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '七、上下文满了我会如何处理', font_size=28, color=TEXT, bold=True)

# Context 管理策略
add_text_box(slide, 0.5, 1.3, 15, 0.5, 'Context 管理策略', font_size=20, color=TEXT, bold=True)
ctx_strategies = [
    ('📦 自动压缩', '系统自动做上下文摘要。之前的对话被压缩成摘要+保留最近几轮的完整内容。Memory文件不受影响。'),
    ('📋 关键信息持久化', '重要的结论写入Memory（/memory命令或我主动建议）。Memory在新会话中自动加载，不受压缩影响。'),
    ('📖 文档始终可读', 'PRD/ARCH/CHECKLIST在文件系统中。我需要时用Read重新读取。不依赖"记得"——依赖"能找到"。'),
]
for i, (title, desc) in enumerate(ctx_strategies):
    y = 2.0 + i * 1.3
    add_rounded_rect(slide, 0.5, y, 15, 1.1, BG_CONTENT, BORDER)
    add_text_box(slide, 0.8, y + 0.15, 14, 0.35, title, font_size=16, color=TEXT, bold=True)
    add_text_box(slide, 0.8, y + 0.55, 14, 0.45, desc, font_size=12, color=TEXT_BODY)

# 防丢策略
add_text_box(slide, 0.5, 5.2, 15, 0.4, '对"上下文丢失"的防御', font_size=20, color=TEXT, bold=True)
defenses = [
    ('你的偏好/约定', '写入 Memory → 新会话自动加载 MEMORY.md 索引'),
    ('当前任务的中间状态', 'git stash / branch / commit 保存代码进度。文档改动直接写入文件'),
    ('重要的讨论结论', '写入 dev_docs/ 对应文档或 Memory'),
    ('项目事实（表数/端点数）', 'gen_arch_stats.py 随时重新生成，不依赖记忆'),
    ('PRD 验收条件', '存在 MD 文件中。Read 即可恢复'),
]
for i, (info, strategy) in enumerate(defenses):
    y = 5.8 + i * 0.5
    add_rounded_rect(slide, 0.5, y, 15, 0.42, BG_CONTENT, BORDER)
    add_text_box(slide, 0.8, y + 0.06, 4, 0.3, info, font_size=12, color=TEXT, bold=True)
    add_text_box(slide, 5.0, y + 0.06, 10, 0.3, strategy, font_size=11, color=TEXT_BODY)

add_slide_number(slide, 8)

# ═══════════════════════ Slide 9: 能力与边界 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '八、我的能力与边界', font_size=28, color=TEXT, bold=True)

# 能做的
add_rounded_rect(slide, 0.5, 1.5, 7.3, 6.8, BG_CONTENT, SUCCESS)
add_text_box(slide, 0.8, 1.7, 6.5, 0.5, '✅ 我能做的', font_size=20, color=SUCCESS, bold=True)
can_items = ['读代码、改代码、写新代码', '读文档、改文档、写新文档',
             '运行脚本（gen_arch_stats.py、vite build、pip install）',
             '搜索代码（Grep/Glob/LSP）', '启动/停止应用（python run.py start/stop）',
             '浏览器操作（导航、点击、截图）', '对照PRD验收条件逐条检查',
             '对照ARCH结构描述验证一致性', '跨模块边界检测（gen_arch_stats.py --check-boundaries）',
             '派发子Agent并行探索/分析', '理解你的意图并提案（不只执行）']
for i, item in enumerate(can_items):
    add_text_box(slide, 0.8, 2.4 + i * 0.52, 6.5, 0.48, f'→ {item}', font_size=12, color=TEXT_BODY)

# 不能做的
add_rounded_rect(slide, 8.3, 1.5, 7.3, 6.8, BG_CONTENT, ERROR)
add_text_box(slide, 8.6, 1.7, 6.5, 0.5, '❌ 我不能做的', font_size=20, color=ERROR, bold=True)
cannot_items = ['在没有Android设备时验证截图流', '在没有ADB环境时真机操作',
                '替代你做最终决策（方案给你，你批）', '自动感知代码运行时bug（需要你反馈症状）',
                '记住跨会话的细节（除非写入Memory）', '访问外网受限页面/API（需权限）']
for i, item in enumerate(cannot_items):
    add_text_box(slide, 8.6, 2.4 + i * 0.9, 6.5, 0.8, f'→ {item}', font_size=12, color=TEXT_BODY)

add_slide_number(slide, 9)

# ═══════════════════════ Slide 10: 协作模式 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '九、你 + 我的协作模式', font_size=28, color=TEXT, bold=True)

# 协作闭环
add_text_box(slide, 0.5, 1.5, 15, 0.5, '完整协作闭环', font_size=22, color=TEXT, bold=True)
collab_flow = [('👤 你说\n"要做什么"', CLAUDE), ('📖 我读\nPRD/ARCH', RGBColor(0x88, 0x9D, 0xF0)),
               ('📐 我\n出方案', RGBColor(0xB3, 0x9E, 0xF3)), ('⏸ 等\n你批', WARNING),
               ('💻 我\n执行', SUCCESS), ('🔍 我\n自检', PRIMARY), ('📋 输出\n结果', CLAUDE)]
for i, (label, color) in enumerate(collab_flow):
    x = 0.3 + i * 2.25
    shape = add_rounded_rect(slide, x, 2.3, 2.0, 1.3, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(13)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 0.5, 4.2, 15, 0.6, '你负责"要做什么"和"审批方案"。我负责"怎么做"和"做完了没有"。',
             font_size=16, color=TEXT, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 0.5, 4.8, 15, 0.6, 'PRD 是你的设计文档，CHECKLIST 是你的验收守门员。我用它们确保代码不偏离你的意图。',
             font_size=14, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# 文档体系总结
add_text_box(slide, 0.5, 5.8, 15, 0.5, '项目配置全貌', font_size=20, color=TEXT, bold=True)
summary_items = [
    ('CLAUDE.md', '路由表 · 身份 · 铁律'),
    ('7 Agents', 'architect / developer / prd-writer / reviewer / tester / test-automator / frontend-evaluator'),
    ('14 Rules', 'frontend / backend / architecture / module-boundaries / security / phone-control / conventions 等'),
    ('6 Hooks', 'check-boundary / check-doc-drift / auto-format / post-turn-build / protect-credentials / session-health'),
    ('12 Skills', 'auto-dev / prd-writer / architecture-review / quality-gate / code-health-check 等'),
    ('55+ Memories', '覆盖：文档管理 / test-runner / CSS / 工作流 / 安全 / 编码习惯'),
]
for i, (item, desc) in enumerate(summary_items):
    y = 6.5 + i * 0.38
    add_text_box(slide, 0.8, y, 2.5, 0.33, item, font_size=12, color=TEXT, bold=True)
    add_text_box(slide, 3.5, y, 12, 0.33, desc, font_size=11, color=TEXT_BODY)

add_slide_number(slide, 10)

# ── 保存 ──
output_path = os.path.join(SAVE_DIR, '02-Claude-Code运行机制.pptx')
prs.save(output_path)
print(f'✅ PPT 2 已保存: {output_path}')
