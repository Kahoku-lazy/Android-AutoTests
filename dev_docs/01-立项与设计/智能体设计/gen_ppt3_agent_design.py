#!/usr/bin/env python3
"""生成 PPT 3：智能体体系设计 (from 智能体体系设计-旧版本-20260716.html)"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# ── 颜色常量 ──
PRIMARY = RGBColor(0x19, 0xC8, 0xB9)
TEXT = RGBColor(0x79, 0x4F, 0x27)
TEXT_BODY = RGBColor(0x72, 0x5D, 0x42)
TEXT_SEC = RGBColor(0x9F, 0x92, 0x7D)
TEXT_MUTED = RGBColor(0x8A, 0x7B, 0x66)
BG = RGBColor(0xF8, 0xF8, 0xF0)
BG_CONTENT = RGBColor(0xF7, 0xF3, 0xDF)
BORDER = RGBColor(0xC4, 0xB8, 0x9E)
SUCCESS = RGBColor(0x6F, 0xBA, 0x2C)
WARNING = RGBColor(0xF5, 0xC3, 0x1C)
ERROR = RGBColor(0xE0, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLAUDE = RGBColor(0xF7, 0xCD, 0x67)
AGENTSCOPE = RGBColor(0xF8, 0xA6, 0xB2)
DEVICE = RGBColor(0x88, 0x9D, 0xF0)
PURPLE = RGBColor(0xB3, 0x9E, 0xF3)

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)

SAVE_DIR = os.path.dirname(os.path.abspath(__file__))


def add_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=14, color=TEXT_BODY,
                 bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = 'Noto Sans SC'
    p.alignment = alignment
    return tf


def add_rounded_rect(slide, left, top, width, height, fill_color=BG_CONTENT, border_color=BORDER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    return shape


def add_slide_number(slide, num):
    add_text_box(slide, 15, 8.4, 0.8, 0.4, str(num), font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.RIGHT)


# ═══════════════════════ Slide 1: 标题页 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_rounded_rect(slide, 1.5, 1.5, 13, 5.5, BG_CONTENT, BORDER)

add_text_box(slide, 2, 2.2, 12, 1.5, '智能体体系设计', font_size=44, color=TEXT, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 2, 3.6, 12, 0.8, 'Claude Code 开发智能体 × AgentScope 产品内 AI 助手 · 路由 · 工作流 · 图解',
             font_size=18, color=TEXT_SEC, alignment=PP_ALIGN.CENTER)

tags = [('Claude Code / Cursor', CLAUDE), ('AgentScope :8000', AGENTSCOPE),
        ('6 角色 · 10 Skill', SUCCESS), ('Stage-Gate 咬合', DEVICE)]
for i, (label, color) in enumerate(tags):
    x = 2.5 + i * 2.8
    shape = add_rounded_rect(slide, x, 5.2, 2.5, 0.5, WHITE, color)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(11)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT_BODY
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 2, 6.5, 12, 0.5, 'Android-AutoTests · 2026-07-10', font_size=14, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ═══════════════════════ Slide 2: 双智能体体系 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '双智能体体系：谁构建、谁运行', font_size=28, color=TEXT, bold=True)

# Claude Code
add_rounded_rect(slide, 0.5, 1.5, 7.3, 3.5, BG_CONTENT, CLAUDE)
add_text_box(slide, 0.8, 1.7, 6.5, 0.5, '🏗️ Claude Code — 做产品的智能体', font_size=18, color=TEXT, bold=True)
cc_items = [('用户', '项目开发者'), ('定位', '写 PRD / 改代码 / 审查 / 验收'),
            ('入口', 'CLAUDE.md 路由'), ('操作', '文件系统 · Bash / Read / Edit'),
            ('产出', '代码变更 + HTML 报告')]
for i, (key, val) in enumerate(cc_items):
    y = 2.4 + i * 0.6
    add_text_box(slide, 1.2, y, 1.5, 0.4, key, font_size=12, color=TEXT_MUTED)
    add_text_box(slide, 2.8, y, 4.5, 0.4, val, font_size=12, color=TEXT_BODY)

# AgentScope
add_rounded_rect(slide, 8.3, 1.5, 7.3, 3.5, BG_CONTENT, AGENTSCOPE)
add_text_box(slide, 8.6, 1.7, 6.5, 0.5, '🤖 AgentScope — 用平台的智能体', font_size=18, color=TEXT, bold=True)
as_items = [('用户', '测试人员'), ('定位', '对话式自动化测试'),
            ('入口', 'Vue ChatView SSE'), ('操作', '25 Tool · Django ORM'),
            ('产出', '用例 + 执行结果 + 报告')]
for i, (key, val) in enumerate(as_items):
    y = 2.4 + i * 0.6
    add_text_box(slide, 9.0, y, 1.5, 0.4, key, font_size=12, color=TEXT_MUTED)
    add_text_box(slide, 10.6, y, 4.5, 0.4, val, font_size=12, color=TEXT_BODY)

# 关系图
add_text_box(slide, 0.5, 5.3, 15, 0.5, '关系图：建造者 → 平台 → 内置助手', font_size=18, color=TEXT, bold=True, alignment=PP_ALIGN.CENTER)
rel_steps = [('Claude Code\n建造者', CLAUDE, '改 .vue / .py / Skill'),
             ('Android-AutoTests\n平台', DEVICE, 'Vue :5173 · Django :8765 · AgentScope :8000'),
             ('AI 助手\n建筑物', AGENTSCOPE, '读写业务 DB')]
for i, (label, color, desc) in enumerate(rel_steps):
    x = 0.5 + i * 5.1
    shape = add_rounded_rect(slide, x, 6.0, 4.8, 1.3, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(14)
    shape.text_frame.paragraphs[0].font.color.rgb = WHITE if color != CLAUDE else TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_text_box(slide, x + 0.2, 7.5, 4.4, 0.4, desc, font_size=11, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

add_text_box(slide, 0.5, 8.3, 15, 0.3, '共享 JWT / Rules 知识；Claude 不碰业务库，AgentScope 不碰源码。',
             font_size=11, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
add_slide_number(slide, 2)

# ═══════════════════════ Slide 3: 六角色路由 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, 'Claude Code · 六角色路由与接力', font_size=28, color=TEXT, bold=True)

# 调度优先级
add_text_box(slide, 0.5, 1.3, 15, 0.5, '调度优先级（CLAUDE.md）', font_size=18, color=TEXT, bold=True)
roles = [('需求？', WARNING), ('prd-writer', CLAUDE), ('复杂？', PURPLE),
         ('architect', DEVICE), ('developer', SUCCESS), ('reviewer', PRIMARY), ('tester', AGENTSCOPE)]
for i, (label, color) in enumerate(roles):
    x = 0.3 + i * 2.2
    shape = add_rounded_rect(slide, x, 2.0, 1.9, 0.7, color if i % 2 == 0 else WHITE, color)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 0.5, 2.9, 15, 0.3, '日常小改可直达 developer（auto-dev）；test-automator 在需要脚本化时上场。',
             font_size=11, color=TEXT_MUTED)

# 泳道
add_text_box(slide, 0.5, 3.5, 15, 0.5, '泳道：一条需求怎么接力', font_size=18, color=TEXT, bold=True)
lanes = [
    ('prd-writer', CLAUDE, '澄清 Why/How/AC → 写子 PRD → Gate 1 心智'),
    ('architect', DEVICE, '影响面 → 契约/协议 → 阶段拆分'),
    ('developer', SUCCESS, 'auto-dev → 编码+编译 → 浏览器验证'),
    ('reviewer', PRIMARY, 'P0–P3 → 安全/契约'),
    ('tester', AGENTSCOPE, '环境探测 → 分层测 → 报告'),
    ('automator', PURPLE, 'DB/API/UI 脚本 → tests/functional/'),
]
for i, (role, color, desc) in enumerate(lanes):
    y = 4.2 + i * 0.7
    shape = add_rounded_rect(slide, 0.5, y, 2.2, 0.55, color, BORDER)
    shape.text_frame.paragraphs[0].text = role
    shape.text_frame.paragraphs[0].font.size = Pt(11)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_text_box(slide, 3.0, y + 0.1, 12, 0.35, desc, font_size=12, color=TEXT_BODY)

add_slide_number(slide, 3)

# ═══════════════════════ Slide 4: Auto-Dev 7阶段 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, 'Claude Code Agent — Auto-Dev 7 阶段工作流', font_size=28, color=TEXT, bold=True)

# 7阶段流程
phases = [
    ('Phase -1\n需求提炼', RGBColor(0xF8, 0xA6, 0xB2), '领域映射→只读探索→\n1-3候选方案'),
    ('Phase 0\n探索（只读）', WARNING, '并行定位文件+依赖\n判定Lite/Standard/Strict'),
    ('Phase 1\n方案（需审批）', RGBColor(0xF8, 0xA6, 0xB2), '目标·文件清单·P0/P1\n未批准不得改代码'),
    ('Phase 2\n编码（自动）', SUCCESS, '逐文件修改→ruff/prettier\n→编译→失败自修≤3次'),
    ('Phase 3\n审查', DEVICE, 'Lite:ruff+语义\nStandard:齿轮1+2\nStrict:全四齿轮'),
    ('Phase 4\n测试（降级）', PURPLE, '5秒环境探测→\n全就绪端到端/\n无设备API契约'),
    ('Phase 5\n交付', SUCCESS, '摘要+审查/测试结果\nStandard/Strict输出HTML'),
]
for i, (name, color, desc) in enumerate(phases):
    x = 0.3 + i * 2.25
    shape = add_rounded_rect(slide, x, 1.5, 2.05, 2.6, color, BORDER)
    shape.text_frame.paragraphs[0].text = name
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = WHITE if color not in [WARNING] else TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_text_box(slide, x + 0.1, 3.0, 1.85, 1.1, desc, font_size=9, color=TEXT_BODY, alignment=PP_ALIGN.CENTER)

# Feedback 旁路
add_rounded_rect(slide, 0.5, 4.4, 15, 0.7, ERROR, BORDER)
tf = add_rounded_rect(slide, 0.5, 4.4, 15, 0.7, ERROR, BORDER)
tf.text_frame.paragraphs[0].text = '🔄 Feedback 旁路：用户反馈 → 症状分类 → 四层下钻 → 定位报告 → 回到 Phase 1 正常流水线'
tf.text_frame.paragraphs[0].font.size = Pt(14)
tf.text_frame.paragraphs[0].font.color.rgb = WHITE
tf.text_frame.paragraphs[0].font.bold = True
tf.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# 阶段详情
add_text_box(slide, 0.5, 5.4, 15, 0.4, '各阶段详情', font_size=18, color=TEXT, bold=True)
phase_details = [
    ('Phase -1', '需求提炼', '领域映射→只读探索→输出1-3个候选方案（文件数/预估时间/推荐）→用户选定后进Phase 0'),
    ('Phase 0', '探索（只读）', '并行定位文件+依赖+关键字；判定Lite/Standard/Strict；初始化workflow-manifest.json'),
    ('Phase 1', '方案（需审批）', '目标·文件清单·P0/P1步骤·风险·验证方式。未批准不得改代码'),
    ('Phase 2', '编码（自动）', '逐文件修改→ruff/prettier→vite build/manage.py check→失败自修最多3次'),
    ('Phase 3', '审查', 'Lite: ruff+语义 · Standard: quality-gate 齿轮1+2 · Strict: 全四齿轮+架构审查'),
    ('Phase 4', '测试（降级）', '5秒环境探测→全就绪端到端/无设备API契约/全不可用静态分析；须浏览器看页面'),
    ('Phase 5', '交付', '摘要+审查/测试结果；Standard/Strict输出HTML到tests/functional/{module}/reports/'),
]
for i, (num, name, desc) in enumerate(phase_details):
    col = i % 2
    row = i // 2
    x = 0.5 + col * 7.7
    y = 6.0 + row * 0.75
    add_rounded_rect(slide, x, y, 7.4, 0.65, BG_CONTENT, BORDER)
    add_text_box(slide, x + 0.2, y + 0.05, 1.8, 0.3, f'{num} {name}', font_size=11, color=TEXT, bold=True)
    add_text_box(slide, x + 2.2, y + 0.05, 5.0, 0.55, desc, font_size=9, color=TEXT_MUTED)

add_slide_number(slide, 4)

# ═══════════════════════ Slide 5: 三档强度 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '三档强度判定（Phase 0 规则引擎）', font_size=28, color=TEXT, bold=True)

# 三档卡片
intensities = [
    ('🪶 Lite', CLAUDE, [
        '1文件 ≤50行 / 修Bug / 文案 / 注解',
        '阶段：探索→方案→编码→Review',
        '审查：终端打印',
        '无 HTML 报告',
    ]),
    ('🛡️ Standard', PRIMARY, [
        '2–3文件 / 50–200行 / 单模块',
        '阶段：…→Review→测试',
        '审查：quality-gate 齿轮 1+2',
        '输出 HTML 报告',
    ]),
    ('🏛️ Strict', PURPLE, [
        '新模块 / 跨模块 / API / >200行',
        '阶段：…→测试→架构审查',
        '审查：四齿轮全开',
        'HTML ×2（质量+架构）',
    ]),
]
for i, (title, color, items) in enumerate(intensities):
    x = 0.5 + i * 5.1
    add_rounded_rect(slide, x, 1.5, 4.8, 3.0, BG_CONTENT, color)
    add_text_box(slide, x + 0.3, 1.7, 4.2, 0.5, title, font_size=20, color=TEXT, bold=True)
    for j, item in enumerate(items):
        add_text_box(slide, x + 0.3, 2.4 + j * 0.5, 4.2, 0.45, f'• {item}', font_size=12, color=TEXT_BODY)

# 决策树
add_text_box(slide, 0.5, 4.8, 15, 0.4, '判定决策树', font_size=18, color=TEXT, bold=True)
tree_steps = [
    ('需求', CLAUDE), ('新模块/API\n跨模块？', WARNING), ('是 →', PURPLE), ('Strict', PURPLE),
]
for i, (label, color) in enumerate(tree_steps):
    x = 0.5 + i * 3.5
    shape = add_rounded_rect(slide, x, 5.4, 3.2, 0.7, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(13)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

tree2 = [('否 · 2–3文件？', WARNING), ('是 →', PRIMARY), ('Standard', PRIMARY), ('否 →', CLAUDE), ('Lite', CLAUDE)]
for i, (label, color) in enumerate(tree2):
    x = 0.5 + i * 3.1
    shape = add_rounded_rect(slide, x, 6.3, 2.8, 0.7, color if i % 2 == 1 else BG_CONTENT, color)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 0.5, 7.3, 15, 0.3, '置信度低 → 对比方案让人选', font_size=12, color=TEXT_MUTED)

add_slide_number(slide, 5)

# ═══════════════════════ Slide 6: 反馈闭环 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '场景：Bug 修复与反馈闭环', font_size=28, color=TEXT, bold=True)

# 反馈流程
fb_steps = [('用户反馈\n"点保存没反应"', RGBColor(0xF8, 0xA6, 0xB2)),
            ('F1 症状分类\n"点按钮没反应"', DEVICE),
            ('F2 四层下钻\n浏览器→前端→API→DB', SUCCESS),
            ('F3 定位报告\n根因+证据+方案', PRIMARY),
            ('回到 Phase 1\n正常开发流水线', WARNING)]
for i, (label, color) in enumerate(fb_steps):
    x = 0.3 + i * 3.15
    shape = add_rounded_rect(slide, x, 1.5, 2.9, 1.3, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(13)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# 四层工具 + 错误反思
add_rounded_rect(slide, 0.5, 3.2, 7.3, 2.8, BG_CONTENT, BORDER)
add_text_box(slide, 0.8, 3.4, 6.5, 0.4, '四层排查工具链', font_size=16, color=TEXT, bold=True)
tools_layers = [
    ('L1', '浏览器 · DevTools → Console / Network'),
    ('L2', '前端编译 · npx vite build / curl :5173/src/...'),
    ('L3', 'API · curl :8765/api/... / tail logs/backend.log'),
    ('L4', '数据库 · manage.py shell + ORM 查询'),
]
for i, (level, desc) in enumerate(tools_layers):
    add_text_box(slide, 0.8, 4.0 + i * 0.55, 6.5, 0.5, f'{level}  {desc}', font_size=11, color=TEXT_BODY)

add_rounded_rect(slide, 8.3, 3.2, 7.3, 2.8, BG_CONTENT, BORDER)
add_text_box(slide, 8.6, 3.4, 6.5, 0.4, '错误反思触发规则', font_size=16, color=TEXT, bold=True)
rules_ref = [
    '→ 第 2 次：提醒用户建议加规则',
    '⚠ 第 3 次：强制写入 .claude/rules/',
    '● 第 4 次+：新规则 + 检查旧规则是否失效',
    '例：catch 静默吞错 ×3 → 写入 frontend.md',
]
for i, r in enumerate(rules_ref):
    add_text_box(slide, 8.6, 4.0 + i * 0.55, 6.5, 0.5, r, font_size=11, color=TEXT_BODY)

# 关键原则
add_text_box(slide, 0.5, 6.5, 15, 0.5, '交付后用户说「不对」→ 不猜、不重开大需求，走 Feedback 专用路径，定位后再回到 Phase 1 正常流水线。',
             font_size=14, color=TEXT_SEC, alignment=PP_ALIGN.CENTER)
add_slide_number(slide, 6)

# ═══════════════════════ Slide 7: Stage-Gate ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '与 Stage-Gate 文档漏斗如何咬合', font_size=28, color=TEXT, bold=True)

gates = [
    ('1', '立项', '人 + 调研扩写', SUCCESS, False),
    ('2', 'PRD', 'prd-writer', PRIMARY, True),
    ('3', '设计', 'architect', RGBColor(0x18, 0x5F, 0xA5), True),
    ('4', '拆分', 'architect · Issues', PURPLE, False),
    ('5', '开发测', 'dev→review→test', RGBColor(0xC7, 0x84, 0x0A), True),
    ('6', '复盘', 'github · Memory', ERROR, False),
]
for i, (num, title, desc, color, has_gate) in enumerate(gates):
    x = 0.5 + i * 2.55
    shape = add_rounded_rect(slide, x, 1.5, 2.3, 2.5, BG_CONTENT, color)
    # Gate number circle
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x + 0.8), Inches(1.7), Inches(0.6), Inches(0.6)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.text_frame.paragraphs[0].text = num
    circle.text_frame.paragraphs[0].font.size = Pt(16)
    circle.text_frame.paragraphs[0].font.color.rgb = WHITE
    circle.text_frame.paragraphs[0].font.bold = True
    circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_text_box(slide, x + 0.15, 2.5, 2.0, 0.4, title, font_size=14, color=TEXT, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.15, 2.9, 2.0, 0.5, desc, font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
    if has_gate:
        gate_shape = add_rounded_rect(slide, x + 0.4, 3.4, 1.4, 0.4, RGBColor(0xFF, 0xF5, 0xF5), ERROR)
        gate_shape.text_frame.paragraphs[0].text = f'Gate {num}'
        gate_shape.text_frame.paragraphs[0].font.size = Pt(9)
        gate_shape.text_frame.paragraphs[0].font.color.rgb = ERROR
        gate_shape.text_frame.paragraphs[0].font.bold = True
        gate_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_text_box(slide, 0.5, 4.5, 15, 0.4, '文档六阶段是产品漏斗；智能体是「谁在每个阶段干活」。Gate 仍由人签发。', font_size=14, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ═══════════════════════ Slide 7 下半部分: Skill/Rules/Memory ═══════════════════════
add_text_box(slide, 0.5, 5.2, 15, 0.5, 'Skill · Rules · Memory 分层', font_size=22, color=TEXT, bold=True)
layers = [
    ('① 入口 · CLAUDE.md / AGENTS.md', CLAUDE, '路由表 · 身份 · 验证铁律 · 项目架构真相。决定「找谁」。'),
    ('② Agents · .claude/agents/', DEVICE, '六角色 system prompt + 工具权限。决定「这个人怎么说话、能用什么工具」。'),
    ('③ Skills · 可触发专项流程', PRIMARY, 'auto-dev · prd-writer · quality-gate · functional-testing · architecture-review …'),
    ('④ Rules · 始终生效的硬约束', ERROR, 'frontend / backend / 防火墙 / JWT / phone-control … 不靠「记得」，靠强制加载。'),
    ('⑤ Memory · 教训沉淀', BORDER, '新模式 / 用户纠正 / 排查教训 → .claude/memory/ + MEMORY.md 索引。'),
]
for i, (title, color, desc) in enumerate(layers):
    y = 5.9 + i * 0.65
    shape = add_rounded_rect(slide, 0.5, y, 15, 0.55, BG_CONTENT, color)
    shape.text_frame.paragraphs[0].text = ''
    add_text_box(slide, 0.8, y + 0.1, 6, 0.35, title, font_size=13, color=TEXT, bold=True)
    add_text_box(slide, 7.0, y + 0.1, 8, 0.35, desc, font_size=11, color=TEXT_MUTED)

add_slide_number(slide, 7)

# ═══════════════════════ Slide 8: 三道防火墙 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '模块边界防火墙（三道）', font_size=28, color=TEXT, bold=True)

firewalls = [
    ('🛡️ 防火墙 #1', SUCCESS, 'service.py 互不 import', [
        '✅ 跨 App import Model', '✅ 跨 App import api.py', '❌ 跨 App import service'
    ]),
    ('🛡️ 防火墙 #2', WARNING, '读放开，写收敛', [
        '✅ 跨 App SELECT 直接 ORM', '❌ 跨 App INSERT/UPDATE/DELETE 必须走 api'
    ]),
    ('🛡️ 防火墙 #3', DEVICE, '外部访问只走 API', [
        'Vue → HTTP → Django API → ORM',
        'AgentScope → Tool → ORM/API（同进程）',
        'Admin → ORM → DB'
    ]),
]
for i, (title, color, lead, items) in enumerate(firewalls):
    x = 0.5 + i * 5.1
    shape = add_rounded_rect(slide, x, 1.5, 4.8, 3.0, BG_CONTENT, color)
    shape.text_frame.paragraphs[0].text = ''
    add_text_box(slide, x + 0.3, 1.7, 4.2, 0.5, title, font_size=18, color=TEXT, bold=True)
    add_text_box(slide, x + 0.3, 2.2, 4.2, 0.4, lead, font_size=13, color=TEXT, bold=True)
    for j, item in enumerate(items):
        add_text_box(slide, x + 0.3, 2.8 + j * 0.5, 4.2, 0.45, item, font_size=12, color=TEXT_BODY)

# API Key 生命周期
add_text_box(slide, 0.5, 4.8, 15, 0.5, '安全规则 · API Key 生命周期', font_size=22, color=TEXT, bold=True)
key_steps = [('用户输入\n明文 Key', SUCCESS), ('encrypt_key()\n加密', PRIMARY),
             ('写入 DB\n仅存密文', WARNING), ('decrypt_key()\n后端使用', DEVICE),
             ('mask_key()\n前端脱敏', RGBColor(0xF8, 0xA6, 0xB2)), ('禁止导出\n明文出站', ERROR)]
for i, (label, color) in enumerate(key_steps):
    x = 0.3 + i * 2.6
    shape = add_rounded_rect(slide, x, 5.5, 2.4, 1.1, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = WHITE if color not in [WARNING, CLAUDE] else TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# 三列规则
rules_cols = [
    ('写入时', SUCCESS, ['必须 encrypt_key()', 'DB值不得以 sk- 开头', '已是mask形态则跳过重写']),
    ('读取时', WARNING, ['列表接口不返回Key', '详情用 mask_key()', '完整Key仅5s揭示窗']),
    ('禁止', ERROR, ['日志打印 Key', 'API响应带完整Key', '迁移/备份导出明文']),
]
for i, (title, color, items) in enumerate(rules_cols):
    x = 0.5 + i * 5.1
    add_rounded_rect(slide, x, 6.9, 4.8, 1.6, BG_CONTENT, color)
    add_text_box(slide, x + 0.3, 7.1, 4.2, 0.4, title, font_size=15, color=TEXT, bold=True)
    for j, item in enumerate(items):
        add_text_box(slide, x + 0.3, 7.6 + j * 0.35, 4.2, 0.3, f'• {item}', font_size=11, color=TEXT_BODY)

add_slide_number(slide, 8)

# ═══════════════════════ Slide 9: AgentScope SOP ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, 'AgentScope Agent — SOP 四阶段（产品内）', font_size=28, color=TEXT, bold=True)

# SOP 流程
sop_steps = [
    ('阶段 1\n需求分析与用例设计', CLAUDE),
    ('阶段 2\n元素准备', PURPLE),
    ('阶段 3\n用例创建与调试', PRIMARY),
    ('阶段 4\n任务执行', SUCCESS),
]
for i, (label, color) in enumerate(sop_steps):
    x = 0.5 + i * 3.85
    shape = add_rounded_rect(slide, x, 1.5, 3.55, 1.3, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(16)
    shape.text_frame.paragraphs[0].font.color.rgb = WHITE if color not in [CLAUDE] else TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# 阶段详情
sop_details = [
    ('阶段 1', '需求分析与用例设计', 'Task* + create_test_sop → case_design；方案等用户确认'),
    ('阶段 2', '元素准备', 'fetch/search_elements → mapping + gaps'),
    ('阶段 3', '用例创建与调试', 'save + debug 到全部 PASS 才能执行'),
    ('阶段 4', '任务执行', '锁设备 → run_test → 报告 → 必须释放'),
]
for i, (num, name, desc) in enumerate(sop_details):
    x = 0.4 + i * 3.88
    add_rounded_rect(slide, x, 3.2, 3.6, 1.3, BG_CONTENT, BORDER)
    add_text_box(slide, x + 0.2, 3.35, 3.2, 0.35, f'{num} {name}', font_size=13, color=TEXT, bold=True)
    add_text_box(slide, x + 0.2, 3.8, 3.2, 0.6, desc, font_size=11, color=TEXT_BODY)

# Leader → SubAgent
add_text_box(slide, 0.5, 4.8, 15, 0.5, 'Leader → 5 SubAgent', font_size=20, color=TEXT, bold=True)
sub_agents = [('Leader', CLAUDE), ('element-\ninspector', PURPLE), ('case-writer', PRIMARY),
              ('device-\noperator', SUCCESS), ('test-\nexecutor', AGENTSCOPE), ('report-\nwriter', WARNING)]
for i, (label, color) in enumerate(sub_agents):
    x = 0.3 + i * 2.6
    shape = add_rounded_rect(slide, x, 5.5, 2.3, 1.1, color, BORDER)
    shape.text_frame.paragraphs[0].text = label
    shape.text_frame.paragraphs[0].font.size = Pt(12)
    shape.text_frame.paragraphs[0].font.color.rgb = TEXT
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

add_slide_number(slide, 9)

# ═══════════════════════ Slide 10: 查阅入口 ═══════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG)
add_text_box(slide, 0.5, 0.3, 15, 0.8, '查阅入口与总结', font_size=28, color=TEXT, bold=True)

# 资源索引
add_text_box(slide, 0.5, 1.3, 15, 0.5, '资源索引（源文件只读）', font_size=20, color=TEXT, bold=True)
refs = [
    ('调度入口', 'CLAUDE.md', '路由 · 铁律 · Skills 索引'),
    ('项目真相', 'AGENTS.md', '架构 · 防火墙 · API/DB'),
    ('Auto-Dev', '.claude/skills/auto-dev/SKILL.md', '7 阶段权威定义'),
    ('角色定义', '.claude/agents/*.md', '六角色 prompt'),
    ('结构索引', '00-智能体/结构总览.md', '目录树与对照表'),
]
for i, (name, path, desc) in enumerate(refs):
    y = 2.0 + i * 0.65
    add_rounded_rect(slide, 0.5, y, 15, 0.52, BG_CONTENT, BORDER)
    add_text_box(slide, 0.8, y + 0.08, 2.5, 0.35, name, font_size=12, color=TEXT, bold=True)
    add_text_box(slide, 3.5, y + 0.08, 5, 0.35, path, font_size=11, color=PRIMARY, bold=True)
    add_text_box(slide, 8.8, y + 0.08, 6, 0.35, desc, font_size=11, color=TEXT_MUTED)

# 总结
add_text_box(slide, 0.5, 3.8, 15, 0.5, '智能体体系总览', font_size=20, color=TEXT, bold=True)
summary_items = [
    ('双体系', 'Claude Code（开发侧） + AgentScope（产品内）互不交叉，共享 JWT / Rules 知识'),
    ('六角色', 'prd-writer → architect → developer → reviewer → tester → test-automator 接力'),
    ('七阶段', 'Phase -1(提炼) → Phase 0(探索) → Phase 1(方案审批) → Phase 2(编码自修) → Phase 3(审查) → Phase 4(测试降级) → Phase 5(交付)'),
    ('三档强度', 'Lite（1文件≤50行）→ Standard（2-3文件50-200行）→ Strict（新模块/跨模块/>200行）'),
    ('三道防火墙', 'service.py互不import | 读放开写收敛 | 外部访问只走API'),
    ('SOP四阶段', '需求分析与用例设计 → 元素准备 → 用例创建与调试 → 任务执行'),
    ('Stage-Gate', '6 阶段文档漏斗：立项→PRD→设计→拆分→开发测→复盘，Gate由人签发'),
]
for i, (title, desc) in enumerate(summary_items):
    y = 4.5 + i * 0.55
    add_rounded_rect(slide, 0.5, y, 15, 0.48, BG_CONTENT, BORDER)
    add_text_box(slide, 0.8, y + 0.08, 2.5, 0.32, title, font_size=12, color=TEXT, bold=True)
    add_text_box(slide, 3.5, y + 0.08, 11.5, 0.32, desc, font_size=11, color=TEXT_BODY)

add_text_box(slide, 0.5, 8.3, 15, 0.3, '视觉参考 agent-workflow-report · animal-island-ui', font_size=10, color=TEXT_MUTED)
add_slide_number(slide, 10)

# ── 保存 ──
output_path = os.path.join(SAVE_DIR, '03-智能体体系设计.pptx')
prs.save(output_path)
print(f'✅ PPT 3 已保存: {output_path}')
