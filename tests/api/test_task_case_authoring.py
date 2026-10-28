"""用例编写任务：提交、文档解析、调度与报告。"""

from __future__ import annotations

import io
import json

import pytest

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from rest_framework.test import APIRequestFactory, force_authenticate

from backend.apps.ai_assistant.models import AIAgent
from backend.apps.task_list import api
from backend.apps.task_list.models import Task, TaskEvent
from backend.apps.task_list.task_reports import build_task_flow
from backend.apps.task_list.views import TaskSubmitAPIView
from backend.models.constants import TaskKind, TaskStatus

pytestmark = [pytest.mark.django_db, pytest.mark.api]


@pytest.fixture
def admin_user(db):
    return get_user_model().objects.create_user(
        username="case-admin", password="pass", is_superuser=True, is_staff=True
    )


@pytest.fixture
def case_agent(admin_user):
    return AIAgent.objects.create(
        owner=admin_user,
        name="Agent-02",
        status="active",
        task_kind=TaskKind.CASE_AUTHORING,
        route_configs={"case_authoring": {"name": "Agent-02"}},
    )


def _submit(user, *, data: dict, files: dict | None = None):
    factory = APIRequestFactory()
    request = factory.post("/api/tasks/submit/", data, format="multipart")
    for key, item in (files or {}).items():
        request.FILES[key] = item
    request.user_id = str(user.id)
    force_authenticate(request, user=user)
    return TaskSubmitAPIView.as_view()(request)


def _md(
    name: str = "prd.md", text: str = "# 6.4 按键\n待机下单击 POWER 开机。"
) -> SimpleUploadedFile:
    return SimpleUploadedFile(name, text.encode("utf-8"), content_type="text/markdown")


def _docx(name: str = "prd.docx") -> SimpleUploadedFile:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("6.4 按键", level=2)
    document.add_paragraph("待机下单击 POWER 开机。")
    document.save(buffer)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def test_submit_case_authoring_task_with_documents(admin_user, case_agent, monkeypatch):
    """Agent-02 任务不要求设备与 YAML，文档落 document 附件。"""
    dispatched = []
    monkeypatch.setattr(api, "dispatch_case_authoring", lambda task: dispatched.append(task.id))

    response = _submit(
        admin_user,
        data={"title": "H717D 家电功能用例", "agent_id": case_agent.id},
        files={"documents": _md()},
    )

    assert response.status_code == 200
    task = Task.objects.get(id=response.data["id"])
    assert task.task_kind == TaskKind.CASE_AUTHORING
    assert task.device_serial == ""
    assert task.agent_id == case_agent.id
    assert task.created_by == str(admin_user.id)

    document = task.file_attachments.get(kind="document")
    assert document.filename == "prd.md"
    assert "待机下单击 POWER 开机" in document.body
    assert dispatched == [task.id]


def test_submit_case_authoring_task_with_docx(admin_user, case_agent, monkeypatch):
    """Word 上传不能因为 MarkItDown 缺可选依赖而失败。"""
    monkeypatch.setattr(api, "dispatch_case_authoring", lambda task: task)
    response = _submit(
        admin_user,
        data={"title": "H717D Word 用例", "agent_id": case_agent.id},
        files={"documents": _docx()},
    )
    assert response.status_code == 200
    task = Task.objects.get(id=response.data["id"])
    document = task.file_attachments.get(kind="document")
    assert document.filename == "prd.docx"
    assert "待机下单击 POWER 开机" in document.body


def test_case_authoring_submit_requires_document(admin_user, case_agent, monkeypatch):
    monkeypatch.setattr(api, "dispatch_case_authoring", lambda task: task)
    response = _submit(
        admin_user,
        data={"title": "没有文档", "agent_id": case_agent.id},
    )
    assert response.status_code == 400
    assert "文档" in str(response.data)


def test_case_authoring_submit_rejects_unsupported_document(admin_user, case_agent, monkeypatch):
    monkeypatch.setattr(api, "dispatch_case_authoring", lambda task: task)
    response = _submit(
        admin_user,
        data={"title": "格式不对", "agent_id": case_agent.id},
        files={"documents": SimpleUploadedFile("notes.txt", b"hello", content_type="text/plain")},
    )
    assert response.status_code == 400
    assert "docx" in str(response.data)


def test_dispatch_case_authoring_starts_thread_without_device(admin_user, case_agent, monkeypatch):
    """用例编写任务没有设备，也能置 running 并起线程。"""
    task = api.create_task(
        case_agent,
        title="H717D",
        goal="",
        task_kind=TaskKind.CASE_AUTHORING,
    )
    monkeypatch.setattr(
        "backend.apps.task_list.task_runner.spawn_task_thread",
        lambda task_id, agent_id: None,
    )
    started = api.dispatch_case_authoring(task)
    assert started is not None
    task.refresh_from_db()
    assert task.status == TaskStatus.RUNNING


def test_build_case_requirement_joins_documents(case_agent):
    task = api.create_task(
        case_agent,
        title="H717D",
        goal="",
        task_kind=TaskKind.CASE_AUTHORING,
    )
    api.add_task_attachment(task, kind="document", filename="a.md", body="甲需求", sort_order=0)
    api.add_task_attachment(task, kind="document", filename="b.md", body="乙需求", sort_order=1)
    text = api.build_case_requirement(task)
    assert "## a.md\n甲需求" in text
    assert "## b.md\n乙需求" in text


def test_case_project_code_from_title():
    assert api.case_project_code(Task(title="H717D 家电用例")) == "H717D"
    assert api.case_project_code(Task(title="家电用例")) == "PROJ"


def test_case_authoring_report_exposes_document_and_conclusion(admin_user, case_agent):
    from django.utils import timezone

    document = "## 一、需求原子表\n\n| ID | 需求 |\n|:--:|--|\n| REQ-001 | 开机 |"
    task = Task.objects.create(
        agent=case_agent,
        title="H717D 用例",
        task_kind=TaskKind.CASE_AUTHORING,
        status=TaskStatus.COMPLETED,
        started_at=timezone.now(),
        finished_at=timezone.now(),
        result=json.dumps(
            {
                "status": "success",
                "summary": "测试用例已生成并通过验收",
                "passed": True,
                "document": document,
                "case_import": {
                    "status": "imported",
                    "project_id": 12,
                    "file_id": 34,
                    "path": "/cases/projects/12/files/34",
                    "case_count": 2,
                    "skipped_count": 0,
                    "message": "",
                },
                "verdicts": [{"passed": True, "message": "通过", "verifier": "case_verifier_role"}],
            },
            ensure_ascii=False,
        ),
    )
    api.add_task_attachment(task, kind="document", filename="prd.md", body="需求", sort_order=0)
    TaskEvent.objects.create(
        task=task,
        seq=1,
        event_type="SOP_STEP_STARTED",
        phase="执行",
        block_type="text",
        payload={"step": "家电测试用例设计与自检", "attempt": 1},
    )
    TaskEvent.objects.create(
        task=task,
        seq=2,
        event_type="ThinkingBlockEndEvent",
        phase="推理",
        block_type="text",
        role_name="case_writer_role",
        model="deepseek-v4-pro",
        payload={"text": "先拆需求原子，再匹配方法"},
    )
    TaskEvent.objects.create(
        task=task,
        seq=3,
        event_type="ThinkingBlockEndEvent",
        phase="推理",
        block_type="text",
        role_name="case_verifier_role",
        model="deepseek-flash",
        payload={"text": "逐条核对覆盖"},
    )
    long_thinking = "完整思考" * 1500
    TaskEvent.objects.create(
        task=task,
        seq=4,
        event_type="ThinkingBlockEndEvent",
        phase="推理",
        block_type="text",
        role_name="case_writer_role",
        model="deepseek-v4-pro",
        payload={"text": long_thinking},
    )

    flow = build_task_flow(str(admin_user.id), str(task.id))
    assert flow is not None
    assert flow["task"]["task_kind"] == TaskKind.CASE_AUTHORING
    assert flow["task"]["device"] == ""
    assert "document" not in flow
    assert flow["case_import"]["status"] == "imported"
    assert flow["case_import"]["path"] == "/cases/projects/12/files/34"
    assert flow["case_import"]["case_count"] == 2
    assert flow["documents"] == [{"filename": "prd.md"}]
    assert flow["conclusion"]["outcome"] == "PASS"
    assert flow["cases"] == []
    assert [item["event_type"] for item in flow["events"]] == [
        "SOP_STEP_STARTED",
        "ThinkingBlockEndEvent",
        "ThinkingBlockEndEvent",
        "ThinkingBlockEndEvent",
    ]
    assert "先拆需求原子" in flow["events"][1]["detail"]
    assert flow["events"][1]["role_kind"] == "case_writer"
    assert flow["events"][2]["role_kind"] == "verifier"
    assert flow["events"][3]["detail"] == long_thinking


def test_case_authoring_report_failed_without_document(admin_user, case_agent):
    task = Task.objects.create(
        agent=case_agent,
        title="H717D 用例",
        task_kind=TaskKind.CASE_AUTHORING,
        status=TaskStatus.FAILED,
        result=json.dumps(
            {
                "status": "fail",
                "summary": "测试用例验收未通过",
                "passed": False,
                "document": "",
                "verdicts": [{"passed": False, "message": "[阻断] 缺少原子", "verifier": "v"}],
            },
            ensure_ascii=False,
        ),
    )
    flow = build_task_flow(str(admin_user.id), str(task.id))
    assert flow is not None
    assert flow["conclusion"]["outcome"] == "FAIL"
    assert "document" not in flow
    assert flow["case_import"] is None
    assert flow["conclusion"]["verdicts"][0]["message"] == "[阻断] 缺少原子"
