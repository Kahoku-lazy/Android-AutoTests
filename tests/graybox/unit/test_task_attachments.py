"""任务附件解析：docx / pdf 在 MarkItDown 缺可选依赖时走本地兜底。"""

from __future__ import annotations

import io

import fitz
import pytest

from docx import Document

from backend.apps.task_list import attachments
from backend.apps.task_list.attachments import parse_document_bytes

pytestmark = [pytest.mark.unit]


class _BoomConverter:
    """模拟 MarkItDown 缺 docx/pdf 可选依赖。"""

    def convert(self, path):
        raise RuntimeError("Missing optional dependency [docx]")


def _docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("6.4 按键", level=2)
    document.add_paragraph("待机下单击 POWER 开机。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "状态"
    table.cell(0, 1).text = "动作"
    table.cell(1, 0).text = "待机"
    table.cell(1, 1).text = "单击 POWER"
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "POWER on, response <= 3s")
    data = document.tobytes()
    document.close()
    return data


@pytest.fixture(autouse=True)
def _force_markitdown_failure(monkeypatch):
    """强制走本地兜底，避免环境里装了 mammoth 后测试路径漂移。"""
    monkeypatch.setattr("markitdown.MarkItDown", lambda: _BoomConverter())
    monkeypatch.setattr(attachments.backend_log, "warning", lambda *args, **kwargs: None)


def test_docx_falls_back_to_python_docx():
    text, name = parse_document_bytes("prd.docx", _docx_bytes())
    assert name == "prd.docx"
    assert "## 6.4 按键" in text
    assert "待机下单击 POWER 开机" in text
    assert "| 状态 | 动作 |" in text
    assert "| 待机 | 单击 POWER |" in text


def test_pdf_falls_back_to_pymupdf():
    text, name = parse_document_bytes("prd.pdf", _pdf_bytes())
    assert name == "prd.pdf"
    assert "POWER on, response <= 3s" in text


def test_docx_fallback_failure_is_user_visible(monkeypatch):
    monkeypatch.setattr(
        attachments, "_fallback_docx", lambda path: (_ for _ in ()).throw(ValueError("坏文档"))
    )
    with pytest.raises(attachments.AttachmentError, match="文档解析失败"):
        parse_document_bytes("prd.docx", _docx_bytes())
